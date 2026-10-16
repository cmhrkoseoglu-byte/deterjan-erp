import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.core.roles import DEPO, LABORATUVAR, SEVKIYAT, YONETIM
from apps.inventory import services as inventory
from apps.inventory.models import Lot, StockMovement
from apps.inventory.tests.factories import stocked_lot, storage_location
from apps.masterdata.tests.factories import MaterialFactory
from apps.returns.tests.factories import finished_product

D = Decimal
URL = "/api/v1/traceability"

pytestmark = pytest.mark.usefixtures("invariants")


@pytest.fixture
def lots(db: None) -> tuple[Lot, Lot]:
    raw = stocked_lot("100", material=MaterialFactory.create(code="HM-X"))
    product = finished_product(code="MM-X")
    lot = inventory.create_lot(
        None, material=product, source=Lot.Source.FILLING, status=Lot.Status.RELEASED
    )
    store = storage_location(code="MAMUL")
    place = raw.balances.get().location
    inventory.post_movement(None, raw.material, raw, place, D(-10), "ISSUE_PRODUCTION", None)
    inventory.link_lots(None, raw, lot, D(10))
    inventory.post_movement(None, product, lot, store, D(40), "PRODUCTION_OUTPUT", None)
    inventory.post_movement(None, product, lot, store, D(-15), StockMovement.Type.SHIPMENT, None)
    return raw, lot


@pytest.mark.parametrize("role", [LABORATUVAR, YONETIM])
def test_trace_endpoints(client_for, lots: tuple[Lot, Lot], role: str) -> None:
    raw, lot = lots
    client = client_for(role)
    backward = client.get(f"{URL}/lots/{lot.pk}/backward/")
    assert backward.status_code == 200
    assert backward.data["children"][0]["lot_no"] == raw.lot_no
    forward = client.get(f"{URL}/lots/{raw.pk}/forward/")
    assert forward.status_code == 200
    child = forward.data["children"][0]
    assert child["on_hand"] == "25.0000"
    assert child["shipped"] == "15.0000"
    recall = client.get(f"{URL}/recall/", {"lot": [raw.pk]})
    assert recall.status_code == 200
    assert recall.data["summary"]["lot_count"] == 2
    assert recall.data["shipments"][0]["quantity"] == "15.0000"


@pytest.mark.parametrize("role", [DEPO, SEVKIYAT])
def test_trace_requires_role(client_for, lots: tuple[Lot, Lot], role: str) -> None:
    client = client_for(role)
    raw, lot = lots
    assert client.get(f"{URL}/lots/{lot.pk}/backward/").status_code == 403
    assert client.get(f"{URL}/recall/", {"lot": [raw.pk]}).status_code == 403
    assert client.get(f"{URL}/recall/pdf/", {"lot": [raw.pk]}).status_code == 403


def test_recall_requires_lot(client_for, db: None) -> None:
    response = client_for(LABORATUVAR).get(f"{URL}/recall/")
    assert response.status_code == 400


def test_recall_exports(client_for, lots: tuple[Lot, Lot]) -> None:
    raw, lot = lots
    client = client_for(YONETIM)
    xlsx = client.get(f"{URL}/recall/xlsx/", {"lot": [raw.pk]})
    assert xlsx.status_code == 200
    assert "attachment" in xlsx["Content-Disposition"]
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == [
        "Özet",
        "Etkilenen lotlar",
        "Stokta kalan",
        "Sevkiyatlar",
        "Müşteriler",
    ]
    lot_numbers = [row[0].value for row in book["Etkilenen lotlar"].iter_rows(min_row=2)]
    assert lot_numbers == [raw.lot_no, lot.lot_no]
    pdf = client.get(f"{URL}/recall/pdf/", {"lot": [raw.pk]})
    assert pdf.status_code == 200
    assert pdf["Content-Type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")
