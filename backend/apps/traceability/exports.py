"""Geri çağırma simülasyonu çıktıları: Excel (openpyxl) ve PDF (reportlab, A4 yatay).

İçerik `selectors.recall_simulation` sonucudur: özet, etkilenen lotlar, stokta kalan, sevkiyatlar
ve müşteriler. Türkçe karakterler için reportlab ile gelen Vera yazı tipi kullanılır (CoA ile
aynı).
"""

import io
from collections.abc import Sequence
from decimal import Decimal
from typing import Any
from xml.sax.saxutils import escape

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.quality.coa import FONT, FONT_BOLD, _register_fonts

NO_CUSTOMER = "Müşteri bilgisi yok"

SUMMARY_LABELS = (
    ("lot_count", "Etkilenen lot"),
    ("finished_lot_count", "Etkilenen mamul lot"),
    ("lots_in_stock", "Stokta olan lot"),
    ("shipped_lot_count", "Sevk edilmiş lot"),
    ("shipment_count", "Sevkiyat belgesi"),
    ("customer_count", "Etkilenen müşteri"),
)
LOT_HEADER = ("Lot no", "Malzeme", "Tür", "Durum", "Stokta", "Sevk edilen", "İade", "Birim")
STOCK_HEADER = ("Lot no", "Malzeme", "Depo", "Lokasyon", "Miktar", "Birim")
SHIPMENT_HEADER = ("Sevkiyat", "Tarih", "Müşteri", "Lot no", "Ürün", "Miktar", "Birim")
CUSTOMER_HEADER = ("Müşteri", "Ürün", "Miktar", "Birim", "Lotlar", "Sevkiyatlar")
TYPE_LABELS = {"RAW": "Hammadde", "PACKAGING": "Ambalaj", "BULK": "Dökme", "FINISHED": "Mamul"}


def _number(value: Decimal) -> str:
    text = f"{value:,.4f}".rstrip("0").rstrip(".")
    return text.replace(",", "_").replace(".", ",").replace("_", ".")


def _when(value: Any) -> str:
    return f"{timezone.localtime(value):%d.%m.%Y %H:%M}" if value else ""


def _customer(row: dict[str, Any]) -> str:
    if row.get("customer_name"):
        return f"{row['customer_code']} — {row['customer_name']}"
    return NO_CUSTOMER


def _selected_text(result: dict[str, Any]) -> str:
    return ", ".join(r["lot_no"] for r in result["lots"] if r["is_selected"])


def lot_rows(result: dict[str, Any]) -> list[list[Any]]:
    return [
        [
            r["lot_no"],
            f"{r['material_code']} {r['material_name']}",
            TYPE_LABELS.get(r["material_type"], r["material_type"]),
            r["status_display"],
            r["on_hand"],
            r["shipped"],
            r["returned"],
            r["unit"],
        ]
        for r in result["lots"]
    ]


def stock_rows(result: dict[str, Any]) -> list[list[Any]]:
    return [
        [
            r["lot_no"],
            f"{r['material_code']} {r['material_name']}",
            r["warehouse_code"],
            r["location_code"],
            r["quantity"],
            r["unit"],
        ]
        for r in result["stock"]
    ]


def shipment_rows(result: dict[str, Any]) -> list[list[Any]]:
    return [
        [
            r["document_number"],
            _when(r["shipped_at"]),
            _customer(r),
            r["lot_no"],
            f"{r['material_code']} {r['material_name']}",
            r["quantity"],
            r["unit"],
        ]
        for r in result["shipments"]
    ]


def customer_rows(result: dict[str, Any]) -> list[list[Any]]:
    return [
        [
            _customer(r),
            f"{r['material_code']} {r['material_name']}",
            r["quantity"],
            r["unit"],
            ", ".join(r["lots"]),
            ", ".join(r["documents"]),
        ]
        for r in result["customers"]
    ]


# --- Excel -----------------------------------------------------------------------------------


def _sheet(workbook: Workbook, title: str, header: Sequence[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(list(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for index, name in enumerate(header, start=1):
        width = max([len(str(name))] + [len(str(r[index - 1])) for r in rows[:500]])
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 60)
    sheet.freeze_panes = "A2"


def recall_xlsx(result: dict[str, Any]) -> bytes:
    """Geri çağırma simülasyonu Excel dosyası (özet + dört sayfa)."""
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Özet"
    summary.append(["Geri çağırma simülasyonu"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Seçilen lotlar", _selected_text(result)])
    summary.append(["Rapor zamanı", _when(timezone.now())])
    summary.append([])
    for key, label in SUMMARY_LABELS:
        summary.append([label, result["summary"][key]])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 60
    _sheet(workbook, "Etkilenen lotlar", LOT_HEADER, lot_rows(result))
    _sheet(workbook, "Stokta kalan", STOCK_HEADER, stock_rows(result))
    _sheet(workbook, "Sevkiyatlar", SHIPMENT_HEADER, shipment_rows(result))
    _sheet(workbook, "Müşteriler", CUSTOMER_HEADER, customer_rows(result))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- PDF -------------------------------------------------------------------------------------


def _table(header: Sequence[str], rows: list[list[Any]], widths: Sequence[float]) -> Table:
    cell = ParagraphStyle("cell", fontName=FONT, fontSize=7.5, leading=9)
    head = ParagraphStyle("head", fontName=FONT_BOLD, fontSize=7.5, leading=9)
    data = [[Paragraph(escape(str(h)), head) for h in header]]
    for row in rows:
        data.append(
            [Paragraph(escape(_number(v) if isinstance(v, Decimal) else str(v)), cell) for v in row]
        )
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


def recall_pdf(result: dict[str, Any]) -> bytes:
    """Geri çağırma simülasyonu PDF'i (A4 yatay)."""
    _register_fonts()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Geri çağırma simülasyonu",
    )
    title = ParagraphStyle("title", fontName=FONT_BOLD, fontSize=14, leading=18)
    heading = ParagraphStyle("heading", fontName=FONT_BOLD, fontSize=10, leading=14)
    text = ParagraphStyle("text", fontName=FONT, fontSize=8.5, leading=11)
    company = getattr(settings, "COMPANY_NAME", "")
    story: list[Any] = [Paragraph("Geri çağırma simülasyonu", title)]
    if company:
        story.append(Paragraph(escape(company), text))
    story.append(Paragraph(f"Seçilen lotlar: {escape(_selected_text(result))}", text))
    story.append(Paragraph(f"Rapor zamanı: {_when(timezone.now())}", text))
    story.append(Spacer(1, 4 * mm))
    story.append(
        _table(
            ["Gösterge", "Değer"],
            [[label, result["summary"][key]] for key, label in SUMMARY_LABELS],
            [60 * mm, 30 * mm],
        )
    )
    sections = (
        ("Müşteriler", CUSTOMER_HEADER, customer_rows(result), (60, 60, 22, 14, 59, 58)),
        ("Sevkiyatlar", SHIPMENT_HEADER, shipment_rows(result), (32, 28, 60, 34, 70, 30, 19)),
        ("Stokta kalan", STOCK_HEADER, stock_rows(result), (40, 95, 30, 50, 37, 21)),
        ("Etkilenen lotlar", LOT_HEADER, lot_rows(result), (34, 80, 22, 24, 30, 30, 30, 23)),
    )
    for name, header, rows, widths in sections:
        story.append(Spacer(1, 5 * mm))
        story.append(Paragraph(f"{name} ({len(rows)})", heading))
        if rows:
            story.append(_table(header, rows, [w * mm for w in widths]))
        else:
            story.append(Paragraph("Kayıt yok.", text))
    doc.build(story)
    return buffer.getvalue()
