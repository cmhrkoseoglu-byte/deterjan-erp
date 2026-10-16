"""İSG raporlarının Excel çıktısı (openpyxl): eksik / eskimiş GBF raporu."""

import io
from collections.abc import Sequence
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.hse.selectors import SdsStatus

Column = tuple[str, str]  # (alan, başlık)

SDS_GAPS: Sequence[Column] = (
    ("material_code", "Malzeme kodu"),
    ("material_name", "Malzeme"),
    ("material_type_display", "Tür"),
    ("is_hazardous", "Tehlikeli madde"),
    ("status_display", "Durum"),
    ("revision_no", "Revizyon no"),
    ("revision_date", "Revizyon tarihi"),
    ("valid_until", "Geçerlilik tarihi"),
)


def _cell(value: Any) -> Any:
    if isinstance(value, bool):
        return "Evet" if value else ""
    if isinstance(value, date):
        return f"{value:%d.%m.%Y}"
    return value


def sds_gaps_xlsx(rows: Sequence[SdsStatus]) -> bytes:
    """Eksik / eskimiş GBF raporu (tek sayfa)."""
    workbook = Workbook()
    ws = workbook.active
    assert ws is not None
    ws.title = "Eksik GBF"
    ws.append([label for _, label in SDS_GAPS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_cell(getattr(row, name)) for name, _ in SDS_GAPS])
    for index, (_, label) in enumerate(SDS_GAPS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = max(14, len(label) + 2)
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
