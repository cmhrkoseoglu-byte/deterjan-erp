"""Açılış stoku içe aktarma (Excel / CSV'den geçiş).

Ana verilerdeki gibi: dosya önce `preview_opening` ile denenir (satır bazlı hata listesi, hiçbir
şey yazılmaz), sonra aynı dosya `run_opening_import` ile onaylanır. Hatalı satır varsa hiçbir
satır yazılmaz. Her satır `services.post_opening_balance` ile yazılır.
"""

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import IO, Any

from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.core.exceptions import BusinessRuleError
from apps.core.models import User
from apps.inventory import services
from apps.inventory.models import Lot
from apps.masterdata.importers import (
    FIRST_DATA_ROW,
    MAX_ROWS,
    CellError,
    Column,
    _header_key,
    _read_rows,
    _text,
    choice_parser,
    code_lookup_parser,
    parse_decimal,
    parse_text,
)
from apps.masterdata.models import Location, Material, UnitOfMeasure

KIND = "opening"


def parse_date(value: Any) -> date:
    """Excel tarihi ya da `15.10.2026`, `15/10/2026`, `2026-10-15` biçiminde metin."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for pattern in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise CellError(f"Tarih bekleniyor (gg.aa.yyyy): “{text}”.")


def parse_lot_no(value: Any) -> str:
    return _text(value).upper()


COLUMNS: tuple[Column, ...] = (
    Column(
        "Malzeme kodu",
        "material",
        code_lookup_parser(Material, "Malzeme"),
        required=True,
        example="HM-LAS",
    ),
    Column("Lot no", "lot_no", parse_lot_no, required=True, example="L2026-114"),
    Column(
        "Lokasyon kodu",
        "location",
        code_lookup_parser(Location, "Lokasyon"),
        required=True,
        example="A-01",
    ),
    Column("Miktar", "quantity", parse_decimal, required=True, example="1250,5"),
    Column(
        "Birim",
        "unit",
        code_lookup_parser(UnitOfMeasure, "Birim"),
        example="KG",
        help="Boşsa malzemenin temel birimi",
    ),
    Column("SKT", "expiry_date", parse_date, example="31.12.2027", help="gg.aa.yyyy"),
    Column("Üretim tarihi", "manufactured_at", parse_date, example="01.01.2026"),
    Column("Tedarikçi lot no", "supplier_lot_no", parse_text, example="S-4471"),
    Column(
        "Durum",
        "status",
        choice_parser(Lot.Status, {"Serbest bırakıldı": "RELEASED"}),
        example="Serbest",
        help="Serbest ya da Karantina (boşsa Serbest)",
    ),
)
TITLE = "Açılış stoku"


@dataclass
class OpeningRow:
    row: int
    code: str
    errors: list[dict[str, str]] = field(default_factory=list)
    data: dict[str, Any] = field(default_factory=dict, repr=False)


@dataclass
class OpeningReport:
    rows: list[OpeningRow]
    file_errors: list[str] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return bool(self.file_errors) or any(row.errors for row in self.rows)

    def as_dict(self) -> dict[str, Any]:
        invalid = sum(1 for row in self.rows if row.errors)
        return {
            "kind": KIND,
            "total": len(self.rows),
            "valid": len(self.rows) - invalid,
            "invalid": invalid,
            "to_create": len(self.rows),
            "to_update": 0,
            "file_errors": self.file_errors,
            "rows": [
                {"row": row.row, "code": row.code, "action": "create", "errors": row.errors}
                for row in self.rows
            ],
        }


class OpeningImportHasErrors(BusinessRuleError):
    def __init__(self, report: OpeningReport) -> None:
        super().__init__(
            "import_has_errors", "Dosyada hatalı satırlar var; hiçbir kayıt aktarılmadı."
        )
        self.report = report


def _parse_table(upload: IO[bytes], filename: str) -> OpeningReport:
    table = _read_rows(upload, filename)
    report = OpeningReport(rows=[])
    if not table:
        report.file_errors.append("Dosya boş.")
        return report
    positions: dict[str, int] = {}
    for index, header in enumerate(table[0]):
        positions.setdefault(_header_key(header), index)
    columns: dict[str, int] = {}
    for column in COLUMNS:
        position = positions.get(_header_key(column.header))
        if position is None:
            if column.required:
                report.file_errors.append(f"“{column.header}” sütunu bulunamadı.")
            continue
        columns[column.field] = position
    if report.file_errors:
        return report
    data_rows = [
        (FIRST_DATA_ROW + offset, row)
        for offset, row in enumerate(table[1:])
        if any(_text(cell) for cell in row)
    ]
    if len(data_rows) > MAX_ROWS:
        report.file_errors.append(f"En fazla {MAX_ROWS} satır aktarılabilir.")
        return report
    if not data_rows:
        report.file_errors.append("Dosyada veri satırı yok.")
        return report
    for row_no, row in data_rows:
        cells = {name: row[index] if index < len(row) else None for name, index in columns.items()}
        result = OpeningRow(row=row_no, code="")
        for column in COLUMNS:
            raw = cells.get(column.field)
            if _text(raw) == "":
                if column.required:
                    result.errors.append({"column": column.header, "message": "Zorunlu alan boş."})
                continue
            try:
                result.data[column.field] = column.parse(raw)
            except CellError as exc:
                result.errors.append({"column": column.header, "message": str(exc)})
        material = result.data.get("material")
        result.code = f"{material.code if material else ''} / {result.data.get('lot_no', '')}"
        report.rows.append(result)
    return report


def _check_file_consistency(report: OpeningReport) -> None:
    """Aynı lot birden çok satırda olabilir (farklı lokasyonlar); SKT ve durumu aynı olmalı,
    aynı lot + lokasyon iki kez yazılamaz."""
    first_seen: dict[tuple[int, str], OpeningRow] = {}
    seen_locations: dict[tuple[int, str, int], int] = {}
    for row in report.rows:
        if row.errors:
            continue
        key = (row.data["material"].pk, row.data["lot_no"])
        place = (*key, row.data["location"].pk)
        if place in seen_locations:
            row.errors.append(
                {
                    "column": "Lokasyon kodu",
                    "message": f"Aynı lot ve lokasyon dosyada tekrar ediyor (satır "
                    f"{seen_locations[place]}).",
                }
            )
            continue
        seen_locations[place] = row.row
        first = first_seen.setdefault(key, row)
        for name, header in (("expiry_date", "SKT"), ("status", "Durum")):
            default = Lot.Status.RELEASED if name == "status" else None
            if first is not row and first.data.get(name, default) != row.data.get(name, default):
                row.errors.append(
                    {
                        "column": header,
                        "message": f"Aynı lotun satır {first.row} ile farklı {header} değeri.",
                    }
                )


def _apply(actor: User | None, report: OpeningReport) -> None:
    """Satırları servisle yazar; iş kuralı hatası satıra eklenir (çağıran geri alır)."""
    for row in report.rows:
        if row.errors:
            continue
        try:
            with transaction.atomic():
                services.post_opening_balance(
                    actor,
                    material=row.data["material"],
                    lot_no=row.data["lot_no"],
                    location=row.data["location"],
                    quantity=Decimal(row.data["quantity"]),
                    unit=row.data.get("unit"),
                    expiry_date=row.data.get("expiry_date"),
                    manufactured_at=row.data.get("manufactured_at"),
                    supplier_lot_no=row.data.get("supplier_lot_no", ""),
                    status=row.data.get("status", Lot.Status.RELEASED),
                )
        except BusinessRuleError as exc:
            row.errors.append({"column": "", "message": exc.message})


def _build(actor: User | None, upload: IO[bytes], filename: str) -> OpeningReport:
    report = _parse_table(upload, filename)
    if report.file_errors:
        return report
    _check_file_consistency(report)
    _apply(actor, report)
    return report


def preview_opening(actor: User | None, upload: IO[bytes], filename: str) -> OpeningReport:
    """Dosyayı doğrular ve her satırı geri alınan bir işlemde dener; hiçbir şey yazılmaz."""
    with transaction.atomic():
        report = _build(actor, upload, filename)
        transaction.set_rollback(True)
    return report


def run_opening_import(actor: User | None, upload: IO[bytes], filename: str) -> OpeningReport:
    """Dosyayı tek işlemde aktarır. Hatalı satır varsa `import_has_errors`; hiçbir satır
    yazılmaz."""
    with transaction.atomic():
        report = _build(actor, upload, filename)
        if report.has_errors:
            transaction.set_rollback(True)
            raise OpeningImportHasErrors(report)
    return report


def build_template() -> bytes:
    """İndirilebilir .xlsx şablonu: başlık + örnek satır, ikinci sayfada sütun açıklamaları."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = TITLE
    bold = Font(bold=True)
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    for index, column in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=column.header)
        cell.font = bold
        if column.required:
            cell.fill = required_fill
        sheet.cell(row=2, column=index, value=column.example)
        sheet.column_dimensions[cell.column_letter].width = max(14, len(column.header) + 4)
    sheet.freeze_panes = "A2"
    notes = workbook.create_sheet("Açıklama")
    for row in (
        ("Sütun", "Zorunlu", "Açıklama"),
        *((c.header, "Evet" if c.required else "", c.help) for c in COLUMNS),
        (),
        ("Her satır bir lotun bir lokasyondaki miktarıdır; lot farklı lokasyonlarda olabilir.",),
        ("Lot sistemde yoksa açılır; miktar açılış hareketiyle stoka yazılır.",),
        ("Örnek satırı silip kendi verinizi girin. Hatalı satır varsa hiçbiri aktarılmaz.",),
    ):
        notes.append(list(row))
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["C"].width = 60
    for cell in notes[1]:
        cell.font = bold
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
