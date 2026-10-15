import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.inventory import services
from apps.inventory.models import GoodsReceipt, Lot, StockAdjustment, StockMovement
from apps.inventory.tests.factories import quarantine_location, stocked_lot, storage_location
from apps.masterdata.models import Material
from apps.masterdata.tests.factories import (
    BulkMaterialFactory,
    MaterialFactory,
    SupplierFactory,
    unit,
)

pytestmark = pytest.mark.usefixtures("invariants")
D = Decimal


def url(name, *args, **query):
    path = reverse(f"inventory:{name}", args=args)
    if query:
        path += "?" + "&".join(f"{key}={value}" for key, value in query.items())
    return path


def xlsx_upload(rows, name="acilis.xlsx"):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = name
    return buffer


# --- yetki -----------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "name", ["lot-list", "movement-list", "balance-list", "goods-receipt-list", "adjustment-list"]
)
def test_read_requires_login_but_no_role(name, client_for):
    assert APIClient().get(url(name)).status_code == 401
    assert client_for().get(url(name)).status_code == 200


@pytest.mark.parametrize("role", ["uretim", "laboratuvar", "planlama", "yonetim"])
def test_stock_writes_require_depo(role, client_for):
    client = client_for(role)
    lot = stocked_lot("5")
    assert client.post(url("goods-receipt-list"), {}, format="json").status_code == 403
    assert client.post(url("adjustment-list"), {}, format="json").status_code == 403
    assert client.post(url("transfer"), {}, format="json").status_code == 403
    movement = StockMovement.objects.get(lot=lot)
    assert client.post(url("movement-reverse", movement.pk), {"reason": "x"}).status_code == 403
    assert client.get(url("opening-template")).status_code == 403


@pytest.mark.parametrize("role", ["depo", "uretim", "planlama", "laboratuvar", "yonetim"])
def test_manual_release_is_admin_only(role, client_for):
    """Karar D-5: elle serbest bırakma yalnızca admin; laboratuvar kalite kararıyla bırakır."""
    lot = stocked_lot("5", status=Lot.Status.QUARANTINE)
    client = client_for(role)
    assert client.post(url("lot-release", lot.pk), {"note": "x"}).status_code == 403
    if role != "laboratuvar":
        assert client.post(url("lot-block", lot.pk), {"note": "x"}).status_code == 403
    lot.refresh_from_db()
    assert lot.status == Lot.Status.QUARANTINE


def test_reject_endpoint_removed(client_for):
    """Faz 2: ret yalnızca kalite kararıyla (`/quality/samples/{id}/decide/`)."""
    lot = stocked_lot("5", status=Lot.Status.QUARANTINE)
    response = client_for("admin").post(f"/api/v1/inventory/lots/{lot.pk}/reject/", {"note": "x"})
    assert response.status_code == 404


def test_admin_can_do_everything(client_for):
    lot = stocked_lot("5", location=quarantine_location(), status=Lot.Status.QUARANTINE)
    admin = client_for("admin")
    no_reason = admin.post(
        url("lot-release", lot.pk), {"to_location": storage_location().pk}, format="json"
    )
    assert no_reason.status_code == 400 and no_reason.json()["code"] == "note_required"
    response = admin.post(
        url("lot-release", lot.pk),
        {"note": "Acil sevk", "to_location": storage_location().pk},
        format="json",
    )
    assert response.status_code == 200, response.json()
    assert response.json()["status_note"] == "Elle serbest bırakma (D-5): Acil sevk"


# --- lot -------------------------------------------------------------------------------------


def test_lot_list_filters_and_detail(client_for):
    q = quarantine_location()
    quarantined = stocked_lot("12.5", location=q, status=Lot.Status.QUARANTINE)
    stocked_lot("3")
    empty = services.create_lot(None, material=MaterialFactory(), source=Lot.Source.PURCHASE)
    client = client_for("depo")

    body = client.get(url("lot-list", status="QUARANTINE")).json()
    assert [row["lot_no"] for row in body["results"]] == [empty.lot_no, quarantined.lot_no]
    assert body["results"][1]["total_quantity"] == "12.5000"
    assert body["results"][1]["allowed_transitions"] == ["BLOCKED", "REJECTED", "RELEASED"]
    with_stock = client.get(url("lot-list", status="QUARANTINE", has_stock="true")).json()
    assert [row["id"] for row in with_stock["results"]] == [quarantined.pk]
    search = client.get(url("lot-list", search=quarantined.lot_no)).json()
    assert search["count"] == 1

    detail = client.get(url("lot-detail", quarantined.pk)).json()
    assert detail["balances"] == [
        {
            "location": q.pk,
            "location_code": q.code,
            "location_name": q.name,
            "location_type": "QUARANTINE",
            "quantity": "12.5000",
        }
    ]
    assert detail["status_display"] == "Karantina"
    assert detail["receipt"] is None


def test_release_and_block_via_api(client_for):
    admin, lab = client_for("admin"), client_for("laboratuvar")
    q, store = quarantine_location(), storage_location()
    lot = stocked_lot("10", location=q, status=Lot.Status.QUARANTINE)

    missing = admin.post(url("lot-release", lot.pk), {"note": "Acil"}, format="json")
    assert missing.status_code == 400
    assert missing.json()["code"] == "release_location_required"

    ok = admin.post(
        url("lot-release", lot.pk), {"note": "Acil", "to_location": store.pk}, format="json"
    )
    assert ok.status_code == 200, ok.json()
    assert ok.json()["status"] == "RELEASED"
    assert ok.json()["balances"][0]["location"] == store.pk

    no_note = lab.post(url("lot-block", lot.pk), {}, format="json")
    assert no_note.status_code == 400 and no_note.json()["code"] == "note_required"
    blocked = lab.post(url("lot-block", lot.pk), {"note": "Şikâyet"}, format="json")
    assert blocked.json()["status"] == "BLOCKED"
    again = lab.post(url("lot-block", lot.pk), {"note": "x"}, format="json")
    assert again.status_code == 400 and again.json()["code"] == "invalid_transition"


def test_lot_trace_endpoint(client_for):
    raw = stocked_lot("10")
    bulk = services.create_lot(None, material=BulkMaterialFactory(), source=Lot.Source.PRODUCTION)
    services.link_lots(None, raw, bulk, D("4"))
    client = client_for()

    backward = client.get(url("lot-trace", bulk.pk)).json()
    assert backward["lot_no"] == bulk.lot_no
    assert backward["children"][0]["lot_no"] == raw.lot_no
    assert backward["children"][0]["quantity"] == "4.0000"
    forward = client.get(url("lot-trace", raw.pk, direction="forward")).json()
    assert forward["children"][0]["lot_no"] == bulk.lot_no
    bad = client.get(url("lot-trace", raw.pk, direction="up"))
    assert bad.status_code == 400


def test_lot_label_pdf(client_for):
    lot = stocked_lot("1")
    response = client_for().get(url("lot-label", lot.pk))
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


# --- hareket, bakiye, transfer ---------------------------------------------------------------


def test_movements_and_reverse(client_for):
    loc = storage_location()
    lot = stocked_lot("10", location=loc)
    scrap = services.post_movement(
        None, lot.material, lot, loc, D("-2"), StockMovement.Type.SCRAP, None
    )
    client = client_for("depo")
    body = client.get(url("movement-list", lot=lot.pk)).json()
    assert [row["movement_type"] for row in body["results"]] == ["SCRAP", "RECEIPT"]
    assert body["results"][0]["is_reversed"] is False

    response = client.post(url("movement-reverse", scrap.pk), {"reason": "Hata"}, format="json")
    assert response.status_code == 200, response.json()
    assert response.json()[0]["quantity"] == "2.0000"
    body = client.get(url("movement-list", lot=lot.pk, movement_type="SCRAP")).json()
    assert body["results"][0]["is_reversed"] is True
    no_reason = client.post(url("movement-reverse", scrap.pk), {"reason": ""}, format="json")
    assert no_reason.status_code == 400
    # Hareketler API'den değiştirilemez ve silinemez.
    assert client.patch(url("movement-detail", scrap.pk), {"note": "x"}).status_code == 405
    assert client.delete(url("movement-detail", scrap.pk)).status_code == 405


def test_balances_and_available(client_for):
    material = MaterialFactory()
    loc = storage_location()
    usable = stocked_lot("10", material=material, location=loc)
    stocked_lot("5", material=material, status=Lot.Status.QUARANTINE)
    client = client_for()
    body = client.get(url("balance-list", material=material.pk)).json()
    assert body["count"] == 2
    released = client.get(url("balance-list", material=material.pk, lot_status="RELEASED")).json()
    assert released["count"] == 1 and released["results"][0]["lot_no"] == usable.lot_no
    available = client.get(url("balance-available", material=material.pk))
    assert available.status_code == 200
    assert [row["lot"] for row in available.json()] == [usable.pk]
    assert client.get(url("balance-available")).status_code == 400


def test_transfer_endpoint(client_for):
    a, b = storage_location(), storage_location()
    lot = stocked_lot("10", location=a)
    client = client_for("depo")
    response = client.post(
        url("transfer"),
        {"lot": lot.pk, "from_location": a.pk, "to_location": b.pk, "quantity": "4"},
        format="json",
    )
    assert response.status_code == 201, response.json()
    assert [row["quantity"] for row in response.json()] == ["-4.0000", "4.0000"]
    too_much = client.post(
        url("transfer"),
        {"lot": lot.pk, "from_location": a.pk, "to_location": b.pk, "quantity": "7"},
        format="json",
    )
    assert too_much.status_code == 400 and too_much.json()["code"] == "negative_stock"


# --- mal kabul -------------------------------------------------------------------------------


def test_goods_receipt_flow(client_for):
    client = client_for("depo")
    supplier = SupplierFactory()
    q, store = quarantine_location(), storage_location()
    las = MaterialFactory(requires_incoming_qc=True)
    bottle = MaterialFactory(
        material_type=Material.Type.PACKAGING, requires_incoming_qc=False, base_unit=unit("ADET")
    )
    payload = {
        "supplier": supplier.pk,
        "receipt_date": "2026-10-15",
        "delivery_note_no": "A-123",
        "quarantine_location": q.pk,
        "storage_location": store.pk,
        "lines": [
            {"material": las.pk, "quantity": "1000", "unit": unit("KG").pk, "supplier_lot_no": "X"},
            {"material": bottle.pk, "quantity": "240", "unit": unit("ADET").pk},
        ],
    }
    created = client.post(url("goods-receipt-list"), payload, format="json")
    assert created.status_code == 201, created.json()
    receipt = created.json()
    assert receipt["number"].startswith("MK-") and receipt["status"] == "DRAFT"
    assert [line["line_no"] for line in receipt["lines"]] == [1, 2]

    first = receipt["lines"][0]
    patched = client.patch(
        url("goods-receipt-detail", receipt["id"]),
        {"lines": [{**first, "quantity": "900"}, payload["lines"][1]]},
        format="json",
    )
    assert patched.status_code == 200, patched.json()
    assert [line["quantity"] for line in patched.json()["lines"]] == ["900.0000", "240.0000"]
    assert [line["line_no"] for line in patched.json()["lines"]] == [1, 3]
    header_only = client.patch(
        url("goods-receipt-detail", receipt["id"]), {"delivery_note_no": "A-124"}, format="json"
    )
    assert header_only.json()["delivery_note_no"] == "A-124"
    assert len(header_only.json()["lines"]) == 2

    posted = client.post(url("goods-receipt-post", receipt["id"]))
    assert posted.status_code == 200, posted.json()
    body = posted.json()
    assert body["status"] == "POSTED" and body["status_display"] == "Kaydedildi"
    assert body["lines"][0]["created_lot_status"] == "QUARANTINE"
    assert body["lines"][1]["created_lot_status"] == "RELEASED"
    lot_detail = client.get(url("lot-detail", body["lines"][0]["created_lot"])).json()
    assert lot_detail["receipt"] == {
        "model": "inventory.goodsreceipt",
        "id": receipt["id"],
        "number": receipt["number"],
    }

    locked = client.patch(url("goods-receipt-detail", receipt["id"]), {"note": "x"}, format="json")
    assert locked.status_code == 400 and locked.json()["code"] == "document_locked"
    # Durum PATCH ile değişmez.
    client.patch(url("goods-receipt-detail", receipt["id"]), {"status": "DRAFT"}, format="json")
    assert GoodsReceipt.objects.get(pk=receipt["id"]).status == "POSTED"

    labels = client.get(url("goods-receipt-labels", receipt["id"]))
    assert labels.status_code == 200 and labels.content.startswith(b"%PDF")
    listing = client.get(url("goods-receipt-list", status="POSTED")).json()
    assert listing["results"][0]["line_count"] == 2


def test_goods_receipt_validation_and_cancel(client_for):
    client = client_for("depo")
    bad = client.post(url("goods-receipt-list"), {"receipt_date": "2026-10-15"}, format="json")
    assert bad.status_code == 400 and "supplier" in bad.json()["errors"]
    created = client.post(
        url("goods-receipt-list"),
        {"supplier": SupplierFactory().pk, "receipt_date": "2026-10-15"},
        format="json",
    ).json()
    empty_post = client.post(url("goods-receipt-post", created["id"]))
    assert empty_post.status_code == 400 and empty_post.json()["code"] == "document_empty"
    no_reason = client.post(url("goods-receipt-cancel", created["id"]), {"reason": ""})
    assert no_reason.status_code == 400 and no_reason.json()["code"] == "cancel_reason_required"
    cancelled = client.post(
        url("goods-receipt-cancel", created["id"]), {"reason": "Mükerrer"}, format="json"
    )
    assert cancelled.json()["status"] == "CANCELLED"
    assert cancelled.json()["cancel_reason"] == "Mükerrer"


# --- sayım düzeltme --------------------------------------------------------------------------


def test_adjustment_flow(client_for):
    client = client_for("depo")
    loc = storage_location()
    lot = stocked_lot("50", location=loc)
    created = client.post(
        url("adjustment-list"),
        {
            "adjustment_date": "2026-10-15",
            "lines": [
                {"lot": lot.pk, "location": loc.pk, "counted_quantity": "48", "reason": "Fire"}
            ],
        },
        format="json",
    )
    assert created.status_code == 201, created.json()
    line = created.json()["lines"][0]
    assert (line["system_quantity"], line["difference"]) == ("50.0000", "-2.0000")
    posted = client.post(url("adjustment-post", created.json()["id"]))
    assert posted.status_code == 200, posted.json()
    assert posted.json()["status"] == "POSTED"
    assert posted.json()["lines"][0]["movement"] is not None
    assert services.current_balance(lot, loc) == D("48")
    assert StockAdjustment.objects.get().posted_by.has_role("depo")


# --- açılış stoku ----------------------------------------------------------------------------

HEADER = ["Malzeme kodu", "Lot no", "Lokasyon kodu", "Miktar", "Birim", "SKT", "Durum"]


def test_opening_import_preview_then_confirm(client_for):
    material = MaterialFactory(code="HM-LAS")
    a, b = storage_location(code="A-01"), storage_location(code="B-01")
    rows = [
        HEADER,
        ["HM-LAS", "l-1", "A-01", "1.250,5", "", "31.12.2027", "Serbest"],
        ["hm-las", "L-1", "B-01", "500", "G", date(2027, 12, 31), ""],
        ["HM-LAS", "L-2", "A-01", 10, "KG", "", "Karantina"],
    ]
    client = client_for("depo")
    preview = client.post(url("opening-import"), {"file": xlsx_upload(rows)}, format="multipart")
    assert preview.status_code == 200, preview.json()
    assert preview.json()["valid"] == 3 and preview.json()["imported"] is False
    assert not Lot.objects.exists()

    confirm = client.post(
        url("opening-import"), {"file": xlsx_upload(rows), "confirm": "true"}, format="multipart"
    )
    assert confirm.status_code == 200, confirm.json()
    lot = Lot.objects.get(lot_no="L-1")
    assert (lot.source, lot.status, lot.expiry_date) == ("OPENING", "RELEASED", date(2027, 12, 31))
    assert services.current_balance(lot, a) == D("1250.5")
    assert services.current_balance(lot, b) == D("0.5")
    assert Lot.objects.get(lot_no="L-2").status == "QUARANTINE"
    assert material.lots.count() == 2


def test_opening_import_errors_write_nothing(client_for):
    MaterialFactory(code="HM-1")
    storage_location(code="A-01")
    rows = [
        HEADER,
        ["HM-1", "L-1", "A-01", "5", "", "31.12.2027", ""],
        ["HM-1", "L-1", "A-01", "5", "", "", ""],  # aynı lot + lokasyon
        ["HM-1", "L-3", "YOK", "abc", "", "32.13.2027", "Reddedildi"],
        ["HM-1", "L-4", "A-01", "5", "ADET", "", ""],  # dönüşüm yok
    ]
    response = client_for("depo").post(
        url("opening-import"), {"file": xlsx_upload(rows), "confirm": "true"}, format="multipart"
    )
    assert response.status_code == 400
    report = response.json()["report"]
    errors = {row["row"]: row["errors"] for row in report["rows"]}
    assert errors[2] == []
    assert errors[3][0]["column"] == "Lokasyon kodu"
    assert {e["column"] for e in errors[4]} == {"Lokasyon kodu", "Miktar", "SKT"}
    assert "dönüşümü tanımlı değil" in errors[5][0]["message"]
    assert not Lot.objects.exists()


def test_opening_import_missing_column_and_template(client_for):
    client = client_for("depo")
    response = client.post(
        url("opening-import"), {"file": xlsx_upload([["Malzeme kodu"]])}, format="multipart"
    )
    assert response.status_code == 200
    assert "“Lot no” sütunu bulunamadı." in response.json()["file_errors"]
    template = client.get(url("opening-template"))
    assert template.status_code == 200
    sheet = load_workbook(io.BytesIO(template.content)).active
    assert [cell.value for cell in sheet[1]][:4] == [
        "Malzeme kodu",
        "Lot no",
        "Lokasyon kodu",
        "Miktar",
    ]


def test_expired_lot_flag(client_for):
    lot = stocked_lot("1", expiry_date=timezone.localdate() - timedelta(days=1))
    body = client_for().get(url("lot-detail", lot.pk)).json()
    assert body["is_expired"] is True
