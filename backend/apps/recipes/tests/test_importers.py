import io
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.core.models import AuditLog
from apps.masterdata.importers import ImportHasErrors
from apps.masterdata.models import Material
from apps.masterdata.tests.factories import BulkMaterialFactory, MaterialFactory, unit
from apps.recipes import importers
from apps.recipes.models import Recipe, RecipeVersion
from apps.recipes.tests.factories import RecipeFactory, RecipeVersionFactory

pytestmark = pytest.mark.django_db

HEADERS = [
    "Reçete kodu",
    "Reçete adı",
    "Ürün kodu",
    "Şarj miktarı",
    "Şarj birimi",
    "Malzeme kodu",
    "Miktar",
    "Birim",
    "Tolerans (%)",
    "Şebeke suyu",
]


def xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = "r.xlsx"
    return buffer


@pytest.fixture
def materials():
    BulkMaterialFactory(code="YM-BUL")
    MaterialFactory(code="HM-SU")
    MaterialFactory(code="HM-LAS")


def recipe_file(code="RC-BUL"):
    return xlsx(
        [
            HEADERS,
            [code, "Bulaşık", "YM-BUL", 1000, "KG", "HM-SU", 700, "KG", None, "Evet"],
            [code, None, None, 1000, "KG", "HM-LAS", "120,5", None, 1, "Hayır"],
        ]
    )


def test_import_creates_recipe_with_draft_version(materials, make_user):
    user = make_user("planlama")

    report = importers.run_import(user, recipe_file(), "r.xlsx").as_dict()

    assert report["invalid"] == 0
    recipe = Recipe.objects.get(code="RC-BUL")
    version = recipe.versions.get()
    assert version.status == RecipeVersion.Status.DRAFT
    assert version.batch_size == Decimal("1000")
    lines = list(version.lines.order_by("sequence"))
    assert [line.material.code for line in lines] == ["HM-SU", "HM-LAS"]
    assert lines[0].is_water and lines[0].tolerance_pct == Decimal("2.00")
    assert lines[1].quantity == Decimal("120.5") and lines[1].unit == unit("KG")
    assert AuditLog.objects.filter(object_id=str(recipe.pk)).exists()


def test_existing_recipe_gets_next_draft_version(materials, make_user):
    recipe = RecipeFactory(code="RC-BUL", product=Material.objects.get(code="YM-BUL"))
    old = RecipeVersionFactory(recipe=recipe)
    old.status = RecipeVersion.Status.APPROVED
    old.save()

    importers.run_import(make_user("planlama"), recipe_file(), "r.xlsx")

    assert list(recipe.versions.values_list("version_no", "status")) == [
        (old.version_no + 1, "DRAFT"),
        (old.version_no, "APPROVED"),
    ]


def test_preview_writes_nothing_and_reports_rule_errors(materials):
    RecipeVersionFactory(recipe=RecipeFactory(code="RC-TASLAK"))
    upload = xlsx(
        [
            HEADERS,
            ["RC-TASLAK", "x", "YM-BUL", 1000, "KG", "HM-SU", 1, None, None, None],
            ["RC-YENI", "Yeni", "HM-LAS", 1000, "KG", "HM-SU", 1, None, None, None],
            ["RC-IKI", "İki", "YM-BUL", 1000, "KG", "YM-BUL", 1, None, None, None],
            ["RC-UC", "Üç", "YM-BUL", 1000, "KG", "HM-SU", 1, None, None, None],
            ["RC-UC", "Üç", "YM-BUL", 500, "KG", "HM-LAS", 1, None, None, None],
        ]
    )
    versions_before = RecipeVersion.objects.count()

    report = importers.preview_import(upload, "r.xlsx").as_dict()

    # Önce hücre ve tutarlılık hataları; RC-UC'nin şarj miktarı satırlar arasında farklı.
    assert report["rows"][4]["errors"][0]["column"] == "Şarj miktarı"
    assert not Recipe.objects.filter(code="RC-YENI").exists()

    upload = xlsx([r for i, r in enumerate(_rows(upload)) if i != 5])
    report = importers.preview_import(upload, "r.xlsx").as_dict()
    messages = [e["message"] for row in report["rows"] for e in row["errors"]]
    assert "taslak sürümü var" in messages[0]
    assert "yarı mamul" in messages[1]  # ürün dökme değil
    assert "kendi satırında" in messages[2]
    assert report["rows"][3]["errors"] == []
    assert RecipeVersion.objects.count() == versions_before
    assert not Recipe.objects.filter(code__in=["RC-YENI", "RC-IKI", "RC-UC"]).exists()


def _rows(upload):
    upload.seek(0)
    return [list(r) for r in load_workbook(upload).active.iter_rows(values_only=True)]


def test_run_import_with_errors_writes_nothing(materials, make_user):
    upload = xlsx(
        [
            HEADERS,
            ["RC-A", "A", "YM-BUL", 1000, "KG", "HM-SU", 1, None, None, None],
            ["RC-B", None, None, 1000, "KG", "HM-SU", 1, None, None, None],
        ]
    )
    with pytest.raises(ImportHasErrors) as exc:
        importers.run_import(make_user("planlama"), upload, "r.xlsx")
    errors = exc.value.report.as_dict()["rows"][1]["errors"]
    assert {e["column"] for e in errors} == {"Reçete adı", "Ürün kodu"}
    assert not Recipe.objects.exists()


def test_template_has_example_recipe():
    sheet = load_workbook(io.BytesIO(importers.build_template())).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == [c.header for c in importers.COLUMNS]
    assert len(rows) == 3


def test_api_import_permissions_and_flow(materials, client_for):
    url = reverse("recipes:import")
    assert client_for("depo").post(url, {"file": recipe_file()}).status_code == 403

    client = client_for("laboratuvar")
    preview = client.post(url, {"file": recipe_file()})
    assert preview.status_code == 200, preview.json()
    assert preview.json()["imported"] is False
    assert not Recipe.objects.exists()

    done = client.post(url, {"file": recipe_file(), "confirm": "true"})
    assert done.json()["imported"] is True
    assert Recipe.objects.get().versions.count() == 1

    again = client.post(url, {"file": recipe_file(), "confirm": "true"})
    assert again.status_code == 400
    assert again.json()["code"] == "import_has_errors"

    template = client.get(reverse("recipes:import-template"))
    assert template.status_code == 200
    assert client_for("depo").get(reverse("recipes:import-template")).status_code == 403
