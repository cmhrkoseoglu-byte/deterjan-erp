import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from apps.core.exceptions import BusinessRuleError
from apps.core.models import AuditLog
from apps.masterdata import importers
from apps.masterdata.models import Material, Supplier
from apps.masterdata.tests.factories import MaterialCategoryFactory, MaterialFactory, unit

pytestmark = pytest.mark.django_db

MATERIAL_HEADERS = ["Kod", "Ad", "Tür", "Temel birim", "Kategori kodu", "Yoğunluk (g/ml)"]


def xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def csv_file(text, encoding="utf-8"):
    return io.BytesIO(text.encode(encoding))


def test_preview_reports_row_errors_and_writes_nothing():
    MaterialCategoryFactory(code="YA")
    upload = xlsx(
        [
            MATERIAL_HEADERS,
            ["hm-las", "LAS", "Hammadde", "kg", "ya", "1,05"],
            ["HM-X", "", "Bilinmeyen", "KG", None, "abc"],
            [None, None, None, None, None, None],  # boş satır atlanır
            ["HM-LAS", "Tekrar", "RAW", "KG", None, None],
        ]
    )

    report = importers.preview_import("materials", upload, "malzeme.xlsx").as_dict()

    assert Material.objects.count() == 0
    assert report["total"] == 3
    assert report["valid"] == 1
    assert report["invalid"] == 2
    first, second, third = report["rows"]
    assert first == {"row": 2, "code": "HM-LAS", "action": "create", "errors": []}
    columns = {error["column"] for error in second["errors"]}
    assert {"Ad", "Tür", "Yoğunluk (g/ml)"} <= columns
    assert third["row"] == 5
    assert "tekrar" in third["errors"][0]["message"]


def test_run_import_creates_and_updates_with_audit(make_user):
    user = make_user("planlama")
    existing = MaterialFactory(code="HM-SLES", name="Eski", density=Decimal("1.0300"))
    upload = xlsx(
        [
            MATERIAL_HEADERS,
            ["HM-SLES", "SLES %70", "Hammadde", "KG", None, None],
            ["AMB-SISE", "Şişe 1L", "Ambalaj", "ADET", None, None],
        ]
    )

    report = importers.run_import(user, "materials", upload, "m.xlsx").as_dict()

    assert (report["to_create"], report["to_update"]) == (1, 1)
    existing.refresh_from_db()
    assert existing.name == "SLES %70"
    assert existing.density == Decimal("1.0300")  # boş hücre mevcut değeri değiştirmez
    bottle = Material.objects.get(code="AMB-SISE")
    assert bottle.material_type == Material.Type.PACKAGING
    assert bottle.base_unit == unit("ADET")
    assert AuditLog.objects.filter(note="Excel içe aktarma").count() == 2


def test_run_import_with_errors_writes_nothing(make_user):
    upload = xlsx(
        [MATERIAL_HEADERS, ["OK1", "Tamam", "RAW", "KG", None, None], ["BAD", "x", "RAW", "YOK"]]
    )

    with pytest.raises(importers.ImportHasErrors) as exc:
        importers.run_import(make_user("planlama"), "materials", upload, "m.xlsx")

    assert exc.value.code == "import_has_errors"
    assert exc.value.report.as_dict()["invalid"] == 1
    assert Material.objects.count() == 0


def test_business_rule_errors_are_reported_per_row():
    MaterialFactory(code="HM-1", base_unit=unit("KG"))
    from apps.recipes.tests.factories import RecipeLineFactory

    RecipeLineFactory(material=Material.objects.get(code="HM-1"))
    upload = xlsx([MATERIAL_HEADERS, ["HM-1", "x", "RAW", "LT", None, None]])

    report = importers.preview_import("materials", upload, "m.xlsx").as_dict()

    assert "temel birimi değiştirilemez" in report["rows"][0]["errors"][0]["message"]


def test_missing_required_column_is_a_file_error():
    upload = xlsx([["Kod", "Ad"], ["X", "Y"]])

    report = importers.preview_import("materials", upload, "m.xlsx").as_dict()

    assert report["file_errors"] == ["“Tür” sütunu bulunamadı.", "“Temel birim” sütunu bulunamadı."]


def test_csv_with_semicolon_and_turkish_encoding(make_user):
    text = "Kod;Ünvan;Vergi no;E-posta\nt001;Örnek Kimya A.Ş.;1234567890;satis@ornek.com\n"

    importers.run_import(make_user("planlama"), "suppliers", csv_file(text, "cp1254"), "t.csv")

    supplier = Supplier.objects.get()
    assert supplier.code == "T001"
    assert supplier.name == "Örnek Kimya A.Ş."


def test_invalid_email_is_reported():
    text = "Kod,Ünvan,E-posta\nM1,Müşteri,gecersiz\n"

    report = importers.preview_import("customers", csv_file(text), "m.csv").as_dict()

    assert report["rows"][0]["errors"][0]["column"] == "E-posta"


def test_unknown_file_type_and_kind_are_rejected():
    with pytest.raises(BusinessRuleError) as exc:
        importers.preview_import("materials", io.BytesIO(b"x"), "m.pdf")
    assert exc.value.code == "import_file_type"

    with pytest.raises(BusinessRuleError) as exc:
        importers.preview_import("lots", io.BytesIO(b"x"), "m.csv")
    assert exc.value.code == "unknown_import_kind"


def test_corrupt_xlsx_is_rejected():
    with pytest.raises(BusinessRuleError) as exc:
        importers.preview_import("materials", io.BytesIO(b"not a zip"), "m.xlsx")
    assert exc.value.code == "import_file_invalid"


@pytest.mark.parametrize(
    ("value", "expected"),
    [("1.234,5", Decimal("1234.5")), ("1,05", Decimal("1.05")), (2.5, Decimal("2.5")), (7, 7)],
)
def test_parse_decimal(value, expected):
    assert importers.parse_decimal(value) == expected


@pytest.mark.parametrize(("value", "expected"), [("Evet", True), ("HAYIR", False), ("1", True)])
def test_parse_bool(value, expected):
    assert importers.parse_bool(value) is expected


@pytest.mark.parametrize("kind", ["materials", "suppliers", "customers"])
def test_template_round_trip_is_valid(kind, make_user):
    """İndirilen şablon (örnek satırıyla) olduğu gibi önizlemeden hatasız geçer."""
    template = importers.build_template(kind)
    workbook = load_workbook(io.BytesIO(template))
    assert workbook.sheetnames[1] == "Açıklama"

    report = importers.preview_import(kind, io.BytesIO(template), "sablon.xlsx").as_dict()

    assert report["file_errors"] == []
    assert report["invalid"] == 0, report["rows"]


@pytest.mark.parametrize(
    ("value", "expected"),
    [("Hammadde", "RAW"), ("raw", "RAW"), ("Yarı mamul", "BULK"), ("DÖKME", "BULK")],
)
def test_material_type_accepts_codes_labels_and_aliases(value, expected):
    parse = importers.SPECS["materials"].columns[2].parse

    assert parse(value) == expected


# --- Geçiş (Faz 5, Adım 5.6): depo, lokasyon, ekipman, birim dönüşümü --------------------------


def test_warehouse_location_equipment_import_in_order(make_user):
    from apps.masterdata.models import Equipment, Location, Warehouse

    user = make_user("planlama")
    importers.run_import(
        user, "warehouses", xlsx([["Kod", "Ad"], ["hmd", "Hammadde deposu"]]), "d.xlsx"
    )
    importers.run_import(
        user,
        "locations",
        xlsx(
            [
                ["Kod", "Ad", "Depo kodu", "Tür"],
                ["A-01", "Raf 1", "HMD", "Depolama"],
                ["TANK-01", "Tank yeri", "HMD", "Tank"],
            ]
        ),
        "l.xlsx",
    )
    report = importers.run_import(
        user,
        "equipment",
        xlsx(
            [
                ["Kod", "Ad", "Tür", "Kapasite", "Kapasite birimi", "Lokasyon kodu"],
                ["T-01", "Tank 1", "Karıştırma tankı", "5000", "KG", "TANK-01"],
                ["H-01", "Hat 1", "Dolum hattı", None, None, None],
            ]
        ),
        "e.xlsx",
    ).as_dict()

    assert report["to_create"] == 2
    assert Warehouse.objects.get(code="HMD").name == "Hammadde deposu"
    assert Location.objects.get(code="TANK-01").location_type == Location.Type.TANK
    tank = Equipment.objects.get(code="T-01")
    assert tank.location.code == "TANK-01"
    assert tank.capacity == Decimal("5000")


def test_location_import_rejects_unknown_warehouse_and_tank_without_location():
    report = importers.preview_import(
        "locations",
        xlsx([["Kod", "Ad", "Depo kodu", "Tür"], ["A-01", "Raf", "YOK", "Depolama"]]),
        "l.xlsx",
    ).as_dict()
    assert report["invalid"] == 1
    assert "Depo bulunamadı" in report["rows"][0]["errors"][0]["message"]

    report = importers.preview_import(
        "equipment",
        xlsx([["Kod", "Ad", "Tür"], ["T-09", "Tank", "Karıştırma tankı"]]),
        "e.xlsx",
    ).as_dict()
    assert report["invalid"] == 1
    assert "lokasyon zorunludur" in report["rows"][0]["errors"][0]["message"]


def test_conversion_import_keys_on_material_and_unit(make_user):
    from apps.masterdata.models import MaterialUnitConversion

    user = make_user("planlama")
    MaterialFactory(code="MM-1L", base_unit=unit("ADET"))
    headers = ["Malzeme kodu", "Birim", "Temel birim karşılığı"]

    first = importers.run_import(
        user, "conversions", xlsx([headers, ["MM-1L", "KOLI", 12]]), "c.xlsx"
    ).as_dict()
    second = importers.preview_import(
        "conversions",
        xlsx([headers, ["mm-1l", "koli", 6], ["MM-1L", "KOLI", 8]]),
        "c.xlsx",
    ).as_dict()

    assert first["rows"][0] == {"row": 2, "code": "MM-1L / KOLI", "action": "create", "errors": []}
    assert second["rows"][0]["action"] == "update"
    assert "tekrar" in second["rows"][1]["errors"][0]["message"]
    importers.run_import(user, "conversions", xlsx([headers, ["MM-1L", "KOLI", 6]]), "c.xlsx")
    assert MaterialUnitConversion.objects.get().factor_to_base == Decimal("6")


@pytest.mark.parametrize("kind", ["warehouses", "locations", "equipment", "conversions"])
def test_new_kinds_have_templates(kind):
    workbook = load_workbook(io.BytesIO(importers.build_template(kind)))
    headers = [cell.value for cell in workbook.worksheets[0][1]]
    assert headers == [c.header for c in importers.get_spec(kind).columns]
