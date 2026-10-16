"""Excel / CSV'den ana veri içe aktarma (Excel'den geçiş).

Akış: dosya önce `preview_import` ile denenir (satır bazlı hata listesi, hiçbir şey yazılmaz),
sonra aynı dosya `run_import` ile onaylanır. Hatalı satır varsa hiçbir satır yazılmaz.
Kodu sistemde olan satır günceller, olmayan yeni kayıt açar. Güncellemede boş bırakılan
isteğe bağlı hücre mevcut değeri değiştirmez.
"""

import csv
import io
import zipfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import IO, Any

from django.core.exceptions import ValidationError
from django.db import models, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from apps.core.exceptions import BusinessRuleError
from apps.core.models import MasterDataModel, User
from apps.masterdata import services
from apps.masterdata.fields import normalize_code
from apps.masterdata.models import (
    Customer,
    Equipment,
    Location,
    Material,
    MaterialCategory,
    MaterialUnitConversion,
    Supplier,
    UnitOfMeasure,
    Warehouse,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
FIRST_DATA_ROW = 2  # 1. satır başlık


class CellError(Exception):
    """Hücre değeri okunamadı (kullanıcıya görünen Türkçe mesaj)."""


# --- hücre çözücüler -------------------------------------------------------------------------


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel 12345 → "12345", "12345.0" değil
    return str(value).strip()


def parse_text(value: Any) -> str:
    return _text(value)


def parse_code(value: Any) -> str:
    return normalize_code(_text(value))


def parse_decimal(value: Any) -> Decimal:
    if isinstance(value, bool):
        raise CellError("Sayı bekleniyor.")
    if isinstance(value, int | Decimal):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = _text(value).replace(" ", "")
    if "," in text and "." in text:  # 1.234,56 (Türkçe) → 1234.56
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise CellError(f"Sayı bekleniyor: “{_text(value)}”.") from exc
    if not number.is_finite():
        raise CellError(f"Sayı bekleniyor: “{_text(value)}”.")
    return number


def parse_int(value: Any) -> int:
    number = parse_decimal(value)
    if number != number.to_integral_value() or number < 0:
        raise CellError(f"Pozitif tam sayı bekleniyor: “{_text(value)}”.")
    return int(number)


_TRUE = {"evet", "e", "var", "1", "true", "doğru", "x"}
_FALSE = {"hayır", "hayir", "h", "yok", "0", "false", "yanlış"}


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = _text(value).replace("I", "ı").replace("İ", "i").lower()
    if text in _TRUE:
        return True
    if text in _FALSE:
        return False
    raise CellError(f"Evet ya da Hayır bekleniyor: “{_text(value)}”.")


def _fold(text: str) -> str:
    return text.replace("I", "ı").replace("İ", "i").lower().strip()


def choice_parser(
    choices: type[models.TextChoices], aliases: dict[str, str] | None = None
) -> Callable[[Any], str]:
    """Kodu (RAW), Türkçe adı (Hammadde) ya da yaygın eş anlamlıyı kabul eden seçim çözücü."""
    lookup = {_fold(value): value for value in choices.values}
    lookup.update({_fold(str(label)): value for value, label in choices.choices})
    lookup.update({_fold(alias): value for alias, value in (aliases or {}).items()})

    def parse(value: Any) -> str:
        found = lookup.get(_fold(_text(value)))
        if found is None:
            allowed = ", ".join(str(label) for label in choices.labels)
            raise CellError(f"Geçersiz değer “{_text(value)}”. Geçerli değerler: {allowed}.")
        return found

    return parse


def code_lookup_parser(model: type[MasterDataModel], label: str) -> Callable[[Any], Any]:
    """Koda göre aktif kayıt bulan çözücü (ör. temel birim, kategori)."""

    def parse(value: Any) -> Any:
        code = parse_code(value)
        obj = model._default_manager.filter(**{"code": code, "is_active": True}).first()
        if obj is None:
            raise CellError(f"{label} bulunamadı ya da pasif: “{code}”.")
        return obj

    return parse


# --- tanımlar --------------------------------------------------------------------------------


@dataclass(frozen=True)
class Column:
    header: str
    field: str
    parse: Callable[[Any], Any]
    required: bool = False
    example: Any = ""
    help: str = ""


@dataclass(frozen=True)
class ImportSpec:
    kind: str
    title: str
    model: type[MasterDataModel]
    columns: tuple[Column, ...]
    # Mevcut kaydı bulan alanlar. Çoğu kayıtta kod; dönüşümde malzeme + birim.
    key_fields: tuple[str, ...] = ("code",)

    @property
    def key_columns(self) -> tuple[Column, ...]:
        return tuple(c for c in self.columns if c.field in self.key_fields)


_PARTNER_COLUMNS = (
    Column("Kod", "code", parse_code, required=True, example="T001", help="Benzersiz kod"),
    Column("Ünvan", "name", parse_text, required=True, example="Örnek Kimya A.Ş."),
    Column("Vergi dairesi", "tax_office", parse_text, example="Kozyatağı"),
    Column("Vergi no", "tax_number", parse_text, example="1234567890"),
    Column("Adres", "address", parse_text, example="İstanbul"),
    Column("Telefon", "phone", parse_text, example="0216 000 00 00"),
    Column("E-posta", "email", parse_text, example="satis@ornek.com"),
    Column("Yetkili kişi", "contact_person", parse_text, example="Ali Veli"),
)

SPECS: dict[str, ImportSpec] = {
    "materials": ImportSpec(
        kind="materials",
        title="Malzemeler",
        model=Material,
        columns=(
            Column("Kod", "code", parse_code, required=True, example="HM-LAS", help="Benzersiz"),
            Column("Ad", "name", parse_text, required=True, example="LAS (Lineer alkil benzen)"),
            Column(
                "Tür",
                "material_type",
                choice_parser(
                    Material.Type,
                    {
                        "Yarı mamul": "BULK",
                        "Yarı mamül": "BULK",
                        "Dökme": "BULK",
                        "Mamül": "FINISHED",
                        "Ürün": "FINISHED",
                    },
                ),
                required=True,
                example="Hammadde",
                help="Hammadde, Ambalaj, Yarı mamul (dökme), Mamul",
            ),
            Column(
                "Temel birim",
                "base_unit",
                code_lookup_parser(UnitOfMeasure, "Birim"),
                required=True,
                example="KG",
                help="Stok bu birimde tutulur: KG, G, LT, ML, ADET, KOLI, PALET",
            ),
            Column(
                "Kategori kodu",
                "category",
                code_lookup_parser(MaterialCategory, "Kategori"),
                example="",
                help="Sistemde tanımlı kategori kodu",
            ),
            Column(
                "Yoğunluk (g/ml)",
                "density",
                parse_decimal,
                example="1,05",
                help="kg ↔ lt dönüşümü için",
            ),
            Column("Raf ömrü (gün)", "shelf_life_days", parse_int, example=730),
            Column("Minimum stok", "min_stock", parse_decimal, example=500, help="Temel birimde"),
            Column(
                "Giriş kontrolü gerekli",
                "requires_incoming_qc",
                parse_bool,
                example="Evet",
                help="Evet / Hayır (boşsa Evet)",
            ),
            Column("Tehlikeli madde", "is_hazardous", parse_bool, example="Hayır"),
            Column("Depolama notu", "storage_notes", parse_text, example="Serin yerde saklayın"),
        ),
    ),
    "suppliers": ImportSpec(
        kind="suppliers", title="Tedarikçiler", model=Supplier, columns=_PARTNER_COLUMNS
    ),
    "customers": ImportSpec(
        kind="customers",
        title="Müşteriler",
        model=Customer,
        columns=tuple(
            Column(c.header, c.field, c.parse, c.required, "M001", c.help)
            if c.field == "code"
            else c
            for c in _PARTNER_COLUMNS
        ),
    ),
    # Geçiş (Faz 5, Adım 5.6): depo → lokasyon → ekipman sırasıyla yüklenir.
    "warehouses": ImportSpec(
        kind="warehouses",
        title="Depolar",
        model=Warehouse,
        columns=(
            Column("Kod", "code", parse_code, required=True, example="HMD", help="Benzersiz"),
            Column("Ad", "name", parse_text, required=True, example="Hammadde deposu"),
        ),
    ),
    "locations": ImportSpec(
        kind="locations",
        title="Lokasyonlar",
        model=Location,
        columns=(
            Column(
                "Kod",
                "code",
                parse_code,
                required=True,
                example="A-01",
                help="Tüm depolarda benzersiz (barkodla okutulur)",
            ),
            Column("Ad", "name", parse_text, required=True, example="A blok 1. raf"),
            Column(
                "Depo kodu",
                "warehouse",
                code_lookup_parser(Warehouse, "Depo"),
                required=True,
                example="HMD",
                help="Önce depoları yükleyin",
            ),
            Column(
                "Tür",
                "location_type",
                choice_parser(Location.Type, {"Raf": "STORAGE", "Depo": "STORAGE"}),
                required=True,
                example="Depolama",
                help="Depolama, Karantina, Üretim, Sevkiyat, Tank",
            ),
        ),
    ),
    "equipment": ImportSpec(
        kind="equipment",
        title="Ekipmanlar",
        model=Equipment,
        columns=(
            Column("Kod", "code", parse_code, required=True, example="T-01", help="Benzersiz"),
            Column("Ad", "name", parse_text, required=True, example="Karıştırma tankı 1"),
            Column(
                "Tür",
                "equipment_type",
                choice_parser(
                    Equipment.Type,
                    {"Tank": "MIXING_TANK", "Hat": "FILLING_LINE", "Kantar": "SCALE"},
                ),
                required=True,
                example="Karıştırma tankı",
                help="Karıştırma tankı, Stok tankı, Dolum hattı, Terazi, Diğer",
            ),
            Column("Kapasite", "capacity", parse_decimal, example=5000),
            Column(
                "Kapasite birimi",
                "capacity_unit",
                code_lookup_parser(UnitOfMeasure, "Birim"),
                example="KG",
                help="Kapasite girildiyse zorunlu",
            ),
            Column(
                "Lokasyon kodu",
                "location",
                code_lookup_parser(Location, "Lokasyon"),
                example="TANK-01",
                help="Tanklar için zorunlu (önce lokasyonları yükleyin)",
            ),
        ),
    ),
    "conversions": ImportSpec(
        kind="conversions",
        title="Birim dönüşümleri",
        model=MaterialUnitConversion,
        key_fields=("material", "unit"),
        columns=(
            Column(
                "Malzeme kodu",
                "material",
                code_lookup_parser(Material, "Malzeme"),
                required=True,
                example="MM-SIVI-1L",
            ),
            Column(
                "Birim",
                "unit",
                code_lookup_parser(UnitOfMeasure, "Birim"),
                required=True,
                example="KOLI",
            ),
            Column(
                "Temel birim karşılığı",
                "factor_to_base",
                parse_decimal,
                required=True,
                example=12,
                help="1 birim kaç temel birim eder (1 KOLI = 12 ADET)",
            ),
        ),
    ),
}


def get_spec(kind: str) -> ImportSpec:
    try:
        return SPECS[kind]
    except KeyError as exc:
        raise BusinessRuleError(
            "unknown_import_kind", f"Tanımsız içe aktarma türü: {kind}"
        ) from exc


# --- dosya okuma -----------------------------------------------------------------------------


def _header_key(text: Any) -> str:
    return _fold(_text(text).rstrip("*").strip())


def _read_rows(upload: IO[bytes], filename: str) -> list[list[Any]]:
    data = upload.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise BusinessRuleError("import_file_too_large", "Dosya en fazla 5 MB olabilir.")
    name = filename.lower()
    if name.endswith(".csv"):
        return _read_csv(data)
    if name.endswith(".xlsx"):
        return _read_xlsx(data)
    raise BusinessRuleError("import_file_type", "Yalnızca .xlsx ya da .csv dosyası yüklenebilir.")


def _read_csv(data: bytes) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "cp1254"):  # Türkçe Excel CSV'yi Windows-1254 kaydeder
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1254 neredeyse her baytı çözer
        raise BusinessRuleError("import_file_encoding", "Dosyanın karakter kodlaması okunamadı.")
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(data: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        raise BusinessRuleError("import_file_invalid", "Excel dosyası okunamadı.") from exc
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


# --- önizleme ve aktarma ---------------------------------------------------------------------


@dataclass
class RowResult:
    row: int
    code: str
    action: str  # "create" | "update"
    errors: list[dict[str, str]] = field(default_factory=list)
    data: dict[str, Any] = field(default_factory=dict, repr=False)
    instance: MasterDataModel | None = field(default=None, repr=False)


@dataclass
class ImportReport:
    kind: str
    rows: list[RowResult]
    file_errors: list[str] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return bool(self.file_errors) or any(row.errors for row in self.rows)

    def as_dict(self) -> dict[str, Any]:
        invalid = sum(1 for row in self.rows if row.errors)
        return {
            "kind": self.kind,
            "total": len(self.rows),
            "valid": len(self.rows) - invalid,
            "invalid": invalid,
            "to_create": sum(1 for row in self.rows if row.action == "create"),
            "to_update": sum(1 for row in self.rows if row.action == "update"),
            "file_errors": self.file_errors,
            "rows": [
                {"row": row.row, "code": row.code, "action": row.action, "errors": row.errors}
                for row in self.rows
            ],
        }


def _validation_messages(exc: ValidationError, spec: ImportSpec) -> list[dict[str, str]]:
    headers = {column.field: column.header for column in spec.columns}
    errors: list[dict[str, str]] = []
    for name, messages in exc.message_dict.items():
        column = headers.get(name, "")
        errors.extend({"column": column, "message": str(message)} for message in messages)
    return errors


def _key_label(value: Any) -> str:
    return str(getattr(value, "code", value))


def _check_row(
    spec: ImportSpec, row_no: int, cells: dict[str, Any], seen: dict[tuple[Any, ...], int]
) -> RowResult:
    result = RowResult(row=row_no, code="", action="create")

    for column in spec.columns:
        raw = cells.get(column.field)
        if _text(raw) == "":
            if column.required:
                result.errors.append({"column": column.header, "message": "Zorunlu alan boş."})
            continue  # boş: yeni kayıtta varsayılan, güncellemede mevcut değer kalır
        try:
            result.data[column.field] = column.parse(raw)
        except CellError as exc:
            result.errors.append({"column": column.header, "message": str(exc)})

    existing = None
    keys = tuple(result.data.get(name) for name in spec.key_fields)
    if all(value not in (None, "") for value in keys):
        result.code = " / ".join(_key_label(value) for value in keys)
        lookup = dict(zip(spec.key_fields, keys, strict=True))
        existing = spec.model._default_manager.filter(**lookup).first()
        result.action = "update" if existing else "create"
        pks = tuple(getattr(value, "pk", value) for value in keys)
        if pks in seen:
            header = spec.key_columns[0].header
            message = f"Kayıt dosyada tekrar ediyor (satır {seen[pks]})."
            result.errors.append({"column": header, "message": message})
        else:
            seen[pks] = row_no
    if result.errors:
        return result

    if existing is not None:
        instance = existing
        for name in spec.key_fields:
            result.data.pop(name, None)
        for name, value in result.data.items():
            setattr(instance, name, value)
    else:
        instance = spec.model(**result.data)
    try:
        services.validate_record(instance)
    except BusinessRuleError as exc:
        result.errors.append({"column": "", "message": exc.message})
    except ValidationError as exc:
        result.errors.extend(_validation_messages(exc, spec))
    result.instance = existing
    return result


def read_table(
    upload: IO[bytes], filename: str, spec_columns: Iterable[Column]
) -> tuple[list[tuple[int, dict[str, Any]]], list[str]]:
    """Dosyayı okur, başlıkları sütun tanımlarıyla eşler; boş satırları atlar.

    Döner: (`(excel satır no, {alan: ham hücre})` listesi, dosya hataları). Başka modüllerin
    içe aktarmaları (ör. reçeteler) da bunu kullanır.
    """
    table = _read_rows(upload, filename)
    if not table:
        return [], ["Dosya boş."]
    positions: dict[str, int] = {}
    for index, header in enumerate(table[0]):
        positions.setdefault(_header_key(header), index)
    columns: dict[str, int] = {}
    errors: list[str] = []
    for column in spec_columns:
        position = positions.get(_header_key(column.header))
        if position is None:
            if column.required:
                errors.append(f"“{column.header}” sütunu bulunamadı.")
            continue
        columns[column.field] = position
    if errors:
        return [], errors

    data_rows = [
        (FIRST_DATA_ROW + offset, row)
        for offset, row in enumerate(table[1:])
        if any(_text(cell) for cell in row)
    ]
    if len(data_rows) > MAX_ROWS:
        return [], [f"En fazla {MAX_ROWS} satır aktarılabilir."]
    if not data_rows:
        return [], ["Dosyada veri satırı yok."]
    return [
        (
            row_no,
            {name: row[index] if index < len(row) else None for name, index in columns.items()},
        )
        for row_no, row in data_rows
    ], []


def _build_report(spec: ImportSpec, upload: IO[bytes], filename: str) -> ImportReport:
    rows, file_errors = read_table(upload, filename, spec.columns)
    report = ImportReport(kind=spec.kind, rows=[], file_errors=file_errors)
    seen: dict[tuple[Any, ...], int] = {}
    for row_no, cells in rows:
        report.rows.append(_check_row(spec, row_no, cells, seen))
    return report


def preview_import(kind: str, upload: IO[bytes], filename: str) -> ImportReport:
    """Dosyayı doğrular, hiçbir şey yazmaz. Satır bazlı hata listesini döner."""
    spec = get_spec(kind)
    # Doğrulama sırasında nesneler değişir; olası yan etkiye karşı geri alınan bir işlemde çalış.
    with transaction.atomic():
        report = _build_report(spec, upload, filename)
        transaction.set_rollback(True)
    return report


def run_import(actor: User, kind: str, upload: IO[bytes], filename: str) -> ImportReport:
    """Dosyayı doğrular ve hatasızsa tüm satırları tek işlemde yazar (denetim günlüklü).

    Hatalı satır varsa `BusinessRuleError("import_has_errors")`; hiçbir satır yazılmaz.
    """
    spec = get_spec(kind)
    with transaction.atomic():
        report = _build_report(spec, upload, filename)
        if report.has_errors:
            raise ImportHasErrors(report)
        for row in report.rows:
            if row.instance is not None:
                fresh = spec.model._default_manager.get(pk=row.instance.pk)
                services.update_record(actor, fresh, row.data, note="Excel içe aktarma")
            else:
                services.create_record(actor, spec.model, row.data, note="Excel içe aktarma")
    return report


class ImportHasErrors(BusinessRuleError):
    """İçe aktarma hatalı satır içeriyor; rapor API yanıtına eklenir."""

    def __init__(self, report: ImportReport) -> None:
        super().__init__(
            "import_has_errors", "Dosyada hatalı satırlar var; hiçbir kayıt aktarılmadı."
        )
        self.report = report


# --- şablon ----------------------------------------------------------------------------------


def build_template(kind: str) -> bytes:
    """İndirilebilir .xlsx şablonu: başlık + örnek satır, ikinci sayfada sütun açıklamaları."""
    spec = get_spec(kind)
    notes = (
        "Kodu sistemde olan satır kaydı günceller, olmayan yeni kayıt açar.",
        "Güncellemede boş bırakılan isteğe bağlı hücre mevcut değeri değiştirmez.",
        "Örnek satırı silip kendi verinizi girin. Hatalı satır varsa hiçbiri aktarılmaz.",
    )
    return build_template_for(spec.title, spec.columns, notes)


def build_template_for(
    title: str,
    columns: Iterable[Column],
    notes: Iterable[str],
    examples: Iterable[Iterable[Any]] | None = None,
) -> bytes:
    """Şablon kitabı: ilk sayfada başlık (zorunlular sarı) + örnek satır(lar), ikinci sayfada
    sütun açıklamaları ve notlar. `examples` verilmezse sütunların `example` değeri kullanılır."""
    columns = tuple(columns)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title[:31]
    header_font = Font(bold=True)
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    rows = [list(row) for row in examples] if examples else [[c.example for c in columns]]
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column.header)
        cell.font = header_font
        if column.required:
            cell.fill = required_fill
        sheet.column_dimensions[cell.column_letter].width = max(14, len(column.header) + 4)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"

    sheet_notes = workbook.create_sheet("Açıklama")
    _write_rows(
        sheet_notes,
        [
            ("Sütun", "Zorunlu", "Açıklama"),
            *((c.header, "Evet" if c.required else "", c.help) for c in columns),
            (),
            *((note,) for note in notes),
        ],
    )
    sheet_notes.column_dimensions["A"].width = 26
    sheet_notes.column_dimensions["C"].width = 60
    for cell in sheet_notes[1]:
        cell.font = header_font

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_rows(sheet: Any, rows: Iterable[tuple[Any, ...]]) -> None:
    for row in rows:
        sheet.append(list(row))
