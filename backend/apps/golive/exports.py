"""Stok karşılaştırması Excel çıktısı ve dosya şablonu."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.golive.models import StockComparison
from apps.golive.services import COMPARISON_COLUMNS
from apps.masterdata.importers import build_template_for

HEADERS = (
    "Malzeme kodu",
    "Malzeme adı",
    "Lot no",
    "Lokasyon",
    "Excel miktarı",
    "Sistem miktarı",
    "Fark (sistem − Excel)",
    "Birim",
    "Not",
)


def comparison_xlsx(comparison: StockComparison) -> bytes:
    """Karşılaştırma sonucu: önce farklı satırlar, sonra eşleşenler."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = comparison.number
    sheet.append([f"Stok karşılaştırması {comparison.number}"])
    sheet.append(
        [
            f"Stok tarihi: {comparison.as_of_date:%d.%m.%Y}",
            f"Düzey: {comparison.get_level_display()}",
            f"Dosya: {comparison.file_name}",
            f"Farklı satır: {comparison.difference_count} / {comparison.line_count}",
        ]
    )
    sheet.append([])
    sheet.append(list(HEADERS))
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    sheet["A1"].font = Font(bold=True, size=13)
    red = PatternFill("solid", fgColor="FDE2E1")
    lines = comparison.lines.select_related("material__base_unit", "location")
    ordered = sorted(lines, key=lambda line: (line.difference == 0, line.material.code))
    for line in ordered:
        note = "" if line.in_file else "Excel'de yok"
        if line.in_file and line.system_quantity == 0 and line.file_quantity:
            note = "Sistemde yok"
        sheet.append(
            [
                line.material.code,
                line.material.name,
                line.lot_no,
                line.location.code if line.location else "",
                line.file_quantity,
                line.system_quantity,
                line.difference,
                line.material.base_unit.code,
                note,
            ]
        )
        if line.difference != 0:
            for cell in sheet[sheet.max_row]:
                cell.fill = red
    for letter, width in zip("ABCDEFGHI", (16, 32, 16, 12, 14, 14, 18, 8, 14), strict=True):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A5"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def comparison_template() -> bytes:
    notes = (
        "Excel'deki stok listenizi bu sütunlarla kaydedin; başka sütunlar yok sayılır.",
        "Lot no tüm satırlarda doluysa lot bazında, lokasyon tüm satırlarda doluysa lokasyon "
        "bazında karşılaştırılır; değilse malzeme toplamı karşılaştırılır.",
        "Aynı malzeme (ve lot/lokasyon) birden çok satırdaysa miktarlar toplanır.",
        "Fark = sistem − Excel. Hatalı satır varsa karşılaştırma yapılmaz.",
    )
    return build_template_for("Stok listesi", COMPARISON_COLUMNS, notes)
