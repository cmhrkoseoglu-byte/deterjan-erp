import io
from datetime import timedelta
from decimal import Decimal
from unittest import mock

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.core import event_names as ev
from apps.core.exceptions import BusinessRuleError
from apps.core.models import AuditLog
from apps.golive import exports, selectors, services
from apps.golive.models import ProblemReport, StockComparison
from apps.inventory import services as inventory
from apps.inventory.tests.factories import stocked_lot, storage_location
from apps.masterdata.importers import ImportHasErrors
from apps.masterdata.tests.factories import MaterialFactory
from apps.workflow.models import Task

pytestmark = [pytest.mark.django_db, pytest.mark.usefixtures("invariants")]

D = Decimal
HEADERS = ["Malzeme kodu", "Lot no", "Lokasyon kodu", "Miktar", "Birim"]


def xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


@pytest.fixture
def stock():
    """İki malzeme: LAS iki lotta (iki lokasyon), SLES tek lotta; kayıtlar servislerle."""
    a, b = storage_location(code="A-01"), storage_location(code="B-01")
    las = MaterialFactory(code="HM-LAS")
    sles = MaterialFactory(code="HM-SLES")
    lot1 = stocked_lot("100", material=las, location=a)
    lot2 = stocked_lot("50", material=las, location=b)
    lot3 = stocked_lot("30", material=sles, location=a)
    return {"a": a, "b": b, "las": las, "sles": sles, "lots": (lot1, lot2, lot3)}


def compare(user, rows, **kwargs):
    return services.create_stock_comparison(
        user, upload=xlsx([HEADERS, *rows]), filename="stok.xlsx", **kwargs
    )


def lines(comparison):
    return {
        (line.material.code, line.lot_no, line.location.code if line.location else ""): (
            line.file_quantity,
            line.system_quantity,
            line.difference,
        )
        for line in comparison.lines.select_related("material", "location")
    }


def test_material_level_comparison_sums_rows_and_lists_missing(stock, make_user):
    user = make_user("depo")
    comparison = compare(
        user,
        [
            ["HM-LAS", None, None, 120, None],
            ["hm-las", None, None, "30", "KG"],  # aynı malzeme: toplanır
        ],
    )

    assert comparison.level == StockComparison.Level.MATERIAL
    assert comparison.number.startswith("SK-")
    assert lines(comparison) == {
        ("HM-LAS", "", ""): (D("150"), D("150"), D("0")),
        ("HM-SLES", "", ""): (D("0"), D("30"), D("30")),  # sistemde var, Excel'de yok
    }
    assert (comparison.line_count, comparison.difference_count) == (2, 1)
    assert not comparison.lines.get(material__code="HM-SLES").in_file
    assert AuditLog.objects.filter(object_id=str(comparison.pk)).exists()


def test_lot_and_location_level_and_only_file_materials(stock, make_user):
    lot1, lot2, _ = stock["lots"]
    comparison = compare(
        make_user("depo"),
        [
            ["HM-LAS", lot1.lot_no, "A-01", 100, None],
            ["HM-LAS", lot2.lot_no, "B-01", 45000, "G"],  # 45 kg: birim dönüşümü
        ],
        only_file_materials=True,
    )

    assert comparison.level == StockComparison.Level.LOT_LOCATION
    assert lines(comparison) == {
        ("HM-LAS", lot1.lot_no, "A-01"): (D("100"), D("100"), D("0")),
        ("HM-LAS", lot2.lot_no, "B-01"): (D("45"), D("50"), D("5")),
    }


def test_mixed_rows_fall_back_to_lower_level(stock, make_user):
    lot1, _, _ = stock["lots"]
    comparison = compare(
        make_user("depo"),
        [["HM-LAS", lot1.lot_no, "A-01", 100, None], ["HM-SLES", None, "A-01", 30, None]],
    )
    assert comparison.level == StockComparison.Level.LOCATION
    assert lines(comparison)[("HM-LAS", "", "A-01")] == (D("100"), D("100"), D("0"))
    assert lines(comparison)[("HM-LAS", "", "B-01")] == (D("0"), D("50"), D("50"))


def test_comparison_uses_balance_at_end_of_as_of_date(make_user):
    las = MaterialFactory(code="HM-LAS")
    past = timezone.now() - timedelta(days=3)
    with mock.patch("django.utils.timezone.now", return_value=past):
        stocked_lot("100", material=las)
    stocked_lot("40", material=las)  # bugün
    yesterday = timezone.localdate() - timedelta(days=1)

    comparison = compare(
        make_user("depo"), [["HM-LAS", None, None, 100, None]], as_of_date=yesterday
    )

    assert lines(comparison)[("HM-LAS", "", "")] == (D("100"), D("100"), D("0"))
    with pytest.raises(BusinessRuleError) as exc:
        compare(
            make_user("depo"),
            [["HM-LAS", None, None, 1, None]],
            as_of_date=timezone.localdate() + timedelta(days=1),
        )
    assert exc.value.code == "future_date"


def test_comparison_rejects_bad_rows_and_writes_nothing(stock, make_user):
    with pytest.raises(ImportHasErrors) as exc:
        compare(
            make_user("depo"),
            [
                ["HM-YOK", None, None, 1, None],
                ["HM-LAS", None, None, -1, None],
                ["HM-LAS", None, None, 1, "LT"],  # yoğunluk yok, dönüşüm yok
                ["HM-SLES", None, None, None, None],
            ],
        )
    rows = exc.value.report.as_dict()["rows"]
    assert [row["errors"][0]["column"] for row in rows] == [
        "Malzeme kodu",
        "Miktar",
        "Birim",
        "Miktar",
    ]
    assert not StockComparison.objects.exists()


def test_export_and_template(stock, make_user):
    comparison = compare(make_user("depo"), [["HM-LAS", None, None, 140, None]])
    sheet = load_workbook(io.BytesIO(exports.comparison_xlsx(comparison))).active
    rows = list(sheet.iter_rows(min_row=5, values_only=True))
    assert rows[0][0] == "HM-LAS" and rows[0][6] == 10  # farklılar önce
    assert rows[1][8] == "Excel'de yok"
    template = load_workbook(io.BytesIO(exports.comparison_template())).active
    assert [c.value for c in template[1]] == HEADERS


def test_system_stock_levels(stock):
    until = services.end_of_day(timezone.localdate())
    total = selectors.system_stock(until, by_lot=False, by_location=False)
    assert total[selectors.StockKey(stock["las"].pk)] == D("150")
    by_loc = selectors.system_stock(until, by_lot=False, by_location=True)
    assert by_loc[selectors.StockKey(stock["las"].pk, "", stock["b"].pk)] == D("50")


# --- hazırlık listesi -------------------------------------------------------------------------


def test_readiness_flags_missing_items(make_user):
    make_user("depo")
    checks = {check["code"]: check for check in selectors.readiness()}

    assert checks["users"]["status"] == "MISSING"
    assert "Planlama" in checks["users"]["items"] and "Depo" not in checks["users"]["items"]
    assert checks["materials_raw"]["status"] == "MISSING"
    assert checks["opening_stock"]["status"] == "MISSING"
    assert checks["go_live_date"]["status"] == "MISSING"
    assert checks["printers"]["status"] == "WARNING"


def test_readiness_reports_opening_stock_and_go_live(make_user):
    user = make_user("depo")
    material = MaterialFactory(code="HM-X")
    inventory.post_opening_balance(
        user, material=material, lot_no="E-1", location=storage_location(), quantity=D("3")
    )
    services.set_go_live_date(user, timezone.localdate() + timedelta(days=3))
    checks = {check["code"]: check for check in selectors.readiness()}
    assert checks["opening_stock"]["status"] == "OK"
    assert checks["opening_stock"]["count"] == 1
    assert checks["materials_raw"]["status"] == "OK"
    assert checks["go_live_date"]["status"] == "OK"


# --- sorun bildir -----------------------------------------------------------------------------


def png(name="ekran.png", size=10):
    return SimpleUploadedFile(name, b"\x89PNG" + b"0" * size, content_type="image/png")


def test_report_problem_opens_admin_task(make_user, django_capture_on_commit_callbacks):
    admin = make_user("admin")
    user = make_user("dolum", first_name="Ayşe", last_name="Yılmaz")

    with django_capture_on_commit_callbacks(execute=True):
        report = services.report_problem(
            user,
            description="  Etiket yazdırılamıyor\nhat 2  ",
            page_url="/dolum/partiler/4",
            page_title="Dolum partisi",
            screenshot=png(),
        )

    assert report.number.startswith("SB-")
    assert report.description == "Etiket yazdırılamıyor\nhat 2"
    assert report.number in report.screenshot.name and report.screenshot.name.endswith(".png")
    task = Task.objects.get()
    assert task.title == f"Sorun bildirimi: {report.number}"
    assert "Ayşe Yılmaz: Etiket yazdırılamıyor hat 2" in task.description
    assert task.assigned_role.name == "admin"
    assert services.can_view_problem(admin, report)
    assert not services.can_view_problem(make_user("dolum"), report)
    assert ev.PROBLEM_REPORTED == "problem_report.created"


def test_report_problem_rules(make_user):
    user = make_user("depo")
    with pytest.raises(BusinessRuleError) as exc:
        services.report_problem(user, description="   ")
    assert exc.value.code == "description_required"
    pdf = SimpleUploadedFile("a.pdf", b"%PDF", content_type="application/pdf")
    with pytest.raises(BusinessRuleError) as exc:
        services.report_problem(user, description="x", screenshot=pdf)
    assert exc.value.code == "file_type_not_allowed"
    with pytest.raises(BusinessRuleError) as exc:
        services.report_problem(user, description="x", screenshot=png(size=9 * 1024 * 1024))
    assert exc.value.code == "file_too_large"
    assert not ProblemReport.objects.exists()
