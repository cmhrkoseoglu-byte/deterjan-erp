"""Puantaj: takvimden varsayılan günler, gün ve aralık girişi, onay ve kilit, tatiller, API."""

from collections.abc import Callable
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from django.core.exceptions import PermissionDenied
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core import events
from apps.core.exceptions import BusinessRuleError
from apps.core.models import AuditLog, User
from apps.personnel import selectors, services
from apps.personnel.models import AttendanceEntry, AttendanceMonth, AttendanceStatus, Employee
from apps.personnel.selectors import AttendanceCell, AttendanceGrid
from apps.personnel.state import attendance_month_states
from apps.personnel.tests.factories import employee

pytestmark = pytest.mark.django_db
BASE = "/api/v1/personnel/attendance"
OTHER_ROLES = ["planlama", "depo", "uretim", "laboratuvar", "dolum", "sevkiyat", "satin_alma"]
S = AttendanceStatus


@pytest.fixture
def manager(make_user: Callable[..., User]) -> User:
    return make_user("yonetim", username="yonetim")


def _previous_month() -> tuple[int, int]:
    """Bitmiş bir ay (onay ancak ay bitince verilir)."""
    last = timezone.localdate().replace(day=1) - timedelta(days=1)
    return last.year, last.month


def _cells(grid: AttendanceGrid, person: Employee) -> dict[date, AttendanceCell]:
    row = next(row for row in grid.rows if row.employee.pk == person.pk)
    return {cell.date: cell for cell in row.cells}


def _code(call: Callable[[], Any]) -> str:
    with pytest.raises(BusinessRuleError) as exc:
        call()
    return exc.value.code


def _workdays(start: date, end: date, *, skip: tuple[date, ...] = ()) -> list[date]:
    return [
        day for day in selectors.days_between(start, end) if day.weekday() != 6 and day not in skip
    ]


# --- takvim ve varsayılan günler -----------------------------------------------------------------


def test_grid_defaults_follow_calendar_and_employment(manager: User) -> None:
    services.add_fixed_holidays(manager, 2026)
    always = employee(manager)
    newcomer = employee(manager, full_name="Can Kaya", hire_date=date(2026, 10, 12))
    leaver = employee(manager, full_name="Ece Ak", hire_date=date(2020, 1, 1))
    services.terminate_employee(manager, leaver, date(2026, 10, 20))
    former = employee(manager, full_name="Eski Personel", hire_date=date(2020, 1, 1))
    services.terminate_employee(manager, former, date(2026, 9, 30))

    grid = selectors.attendance_grid(2026, 10)
    assert [row.employee.full_name for row in grid.rows] == ["Ayşe Demir", "Can Kaya", "Ece Ak"]
    assert grid.status == AttendanceMonth.Status.OPEN and grid.record is None
    for day, cell in _cells(grid, always).items():
        if day == date(2026, 10, 29):
            expected = S.PUBLIC_HOLIDAY
        elif day.weekday() == 6:
            expected = S.WEEKLY_REST
        else:
            expected = S.WORKED  # 28 Ekim yarım gün: çalıştı açılır
        assert (cell.status, cell.explicit, cell.editable) == (expected, False, True)
    arife = next(day for day in grid.days if day.date == date(2026, 10, 28))
    assert arife.half_day and arife.holiday_name == "Cumhuriyet Bayramı arifesi"
    assert _cells(grid, newcomer)[date(2026, 10, 11)].status == ""
    assert not _cells(grid, newcomer)[date(2026, 10, 11)].editable
    assert _cells(grid, newcomer)[date(2026, 10, 12)].status != ""
    assert _cells(grid, leaver)[date(2026, 10, 20)].status != ""
    assert _cells(grid, leaver)[date(2026, 10, 21)].status == ""
    assert len(selectors.attendance_grid(2026, 10, department="Dolum").rows) == 3
    assert selectors.attendance_grid(2026, 10, department="Depo").rows == []


def test_rest_weekdays_setting(manager: User) -> None:
    card = employee(manager)
    services.update_attendance_settings(manager, [6, 5, 5])
    assert services.get_attendance_settings().rest_weekdays == [5, 6]
    cells = _cells(selectors.attendance_grid(2026, 10), card)
    saturday = next(day for day in cells if day.weekday() == 5)
    assert cells[saturday].status == S.WEEKLY_REST
    invalid: list[list[int]] = [[], [0, 1, 2, 3, 4, 5, 6], [7]]
    for days in invalid:
        with pytest.raises(BusinessRuleError) as exc:
            services.update_attendance_settings(manager, days)
        assert exc.value.code == "invalid_rest_weekdays"


def test_holidays(manager: User) -> None:
    holiday = services.add_holiday(
        manager, {"date": date(2026, 3, 20), "name": "  Ramazan   Bayramı 1. gün "}
    )
    assert holiday.name == "Ramazan Bayramı 1. gün" and not holiday.half_day
    duplicate = {"date": date(2026, 3, 20), "name": "Başka"}
    assert _code(lambda: services.add_holiday(manager, duplicate)) == "holiday_exists"
    blank = {"date": date(2026, 3, 21), "name": " "}
    assert _code(lambda: services.add_holiday(manager, blank)) == "name_required"
    assert services.update_holiday(manager, holiday, {"half_day": True}).half_day
    removed = services.remove_holiday(manager, holiday)
    assert not removed.is_active
    services.add_holiday(manager, {"date": date(2026, 3, 20), "name": "Ramazan Bayramı"})

    created = services.add_fixed_holidays(manager, 2027)
    assert [(h.date.month, h.date.day) for h in created] == [
        (1, 1),
        (4, 23),
        (5, 1),
        (5, 19),
        (7, 15),
        (8, 30),
        (10, 28),
        (10, 29),
    ]
    assert [h.half_day for h in created].count(True) == 1
    assert services.add_fixed_holidays(manager, 2027) == []  # tekrar çalıştırmak güvenli
    assert _code(lambda: services.add_fixed_holidays(manager, 1999)) == "invalid_year"


# --- gün ve aralık girişi ------------------------------------------------------------------------


def test_set_day_rules(manager: User) -> None:
    year, month = _previous_month()
    card = employee(manager, hire_date=date(year, month, 2))
    day = date(year, month, 10)
    entry = services.set_attendance_day(
        manager, card, day, status=S.WORKED, overtime_hours=Decimal("2.5"), note=" Sipariş  "
    )
    assert entry.overtime_hours == Decimal("2.5") and entry.note == "Sipariş"
    assert entry.month.number.startswith("PU-") and entry.month.status == "OPEN"
    log = AuditLog.objects.get(
        object_id=entry.pk, action=AuditLog.Action.CREATE, content_type__model="attendanceentry"
    )
    assert log.changes["status"] == [None, "WORKED"]
    again = services.set_attendance_day(manager, card, day, status=S.SICK_LEAVE)
    assert again.pk == entry.pk and again.overtime_hours == 0
    assert AttendanceMonth.objects.count() == 1

    def write(when: date = day, **kwargs: Any) -> Callable[[], Any]:
        return lambda: services.set_attendance_day(manager, card, when, **kwargs)

    overtime = Decimal(1)
    assert _code(write(status=S.ANNUAL_LEAVE, overtime_hours=overtime)) == "overtime_not_allowed"
    assert _code(write(status=S.WORKED, overtime_hours=Decimal(25))) == "invalid_overtime"
    assert _code(write(status="HOLIDAY")) == "invalid_status"
    assert _code(write(date(year, month, 1), status=S.WORKED)) == "not_employed"
    future = timezone.localdate() + timedelta(days=3)
    assert _code(write(future, status=S.ABSENT)) == "future_day"
    assert _code(write(future, status=S.WORKED, overtime_hours=overtime)) == "future_day"
    # İzin ve rapor önceden girilebilir.
    assert services.set_attendance_day(manager, card, future, status=S.ANNUAL_LEAVE).pk


def test_overtime_by_day_type_and_totals(manager: User) -> None:
    year, month = _previous_month()
    first, last = selectors.month_bounds(year, month)
    holiday = date(year, month, 15)
    services.add_holiday(manager, {"date": holiday, "name": "Bayram"})
    card = employee(manager)
    sunday = next(day for day in selectors.days_between(first, last) if day.weekday() == 6)
    workdays = _workdays(first, last, skip=(holiday,))
    services.set_attendance_day(
        manager, card, sunday, status=S.WEEKLY_REST, overtime_hours=Decimal(8)
    )
    services.set_attendance_day(
        manager, card, holiday, status=S.PUBLIC_HOLIDAY, overtime_hours=Decimal(4)
    )
    services.set_attendance_day(
        manager, card, workdays[0], status=S.WORKED, overtime_hours=Decimal("2.5")
    )
    services.set_attendance_day(manager, card, workdays[1], status=S.SICK_LEAVE)
    services.set_attendance_day(manager, card, workdays[2], status=S.UNPAID_LEAVE)
    services.set_attendance_day(manager, card, workdays[3], status=S.ABSENT)

    totals = selectors.attendance_grid(year, month).rows[0].totals
    assert totals.overtime == {
        "WORKED": Decimal("2.5"),
        "WEEKLY_REST": Decimal(8),
        "PUBLIC_HOLIDAY": Decimal(4),
    }
    assert totals.overtime_total == Decimal("14.5")
    assert totals.days[S.SICK_LEAVE] == 1 and totals.missing_days == 3
    assert totals.days[S.PUBLIC_HOLIDAY] == 1
    assert sum(totals.days.values()) == (last - first).days + 1


def test_range_rules(manager: User) -> None:
    year, month = _previous_month()
    first, last = selectors.month_bounds(year, month)
    holiday = date(year, month, 15)
    services.add_holiday(manager, {"date": holiday, "name": "Bayram"})
    filler = employee(manager, full_name="Ayşe Demir", department="Dolum")
    keeper = employee(manager, full_name="Can Kaya", department="Depo")

    # Yıllık izin hafta tatili ve resmi tatil günlerine yazılmaz (İş Kanunu m.56).
    count = services.set_attendance_range(
        manager, start=first, end=last, status=S.ANNUAL_LEAVE, employee=filler
    )
    workdays = _workdays(first, last, skip=(holiday,))
    assert count == len(workdays)
    cells = _cells(selectors.attendance_grid(year, month), filler)
    assert {cells[day].status for day in workdays} == {S.ANNUAL_LEAVE}
    assert cells[holiday].status == S.PUBLIC_HOLIDAY and not cells[holiday].explicit

    # Rapor her takvim gününe yazılır; birim süzgeci yalnızca o birimi değiştirir.
    week = first + timedelta(days=6)
    count = services.set_attendance_range(
        manager, start=first, end=week, status=S.SICK_LEAVE, department="Depo"
    )
    assert count == 7
    grid = selectors.attendance_grid(year, month)
    assert {_cells(grid, keeper)[day].status for day in selectors.days_between(first, week)} == {
        S.SICK_LEAVE
    }
    assert _cells(grid, filler)[first].status != S.SICK_LEAVE

    # Herkes: ücretsiz izin; aynı aralığı tekrar yazmak değişiklik sayılmaz.
    assert services.set_attendance_range(manager, start=last, end=last, status=S.UNPAID_LEAVE) == 2
    assert services.set_attendance_range(manager, start=last, end=last, status=S.UNPAID_LEAVE) == 0

    # İzne çevrilen günün fazla mesaisi sıfırlanır.
    day = workdays[-2]
    services.set_attendance_day(manager, keeper, day, status=S.WORKED, overtime_hours=Decimal(2))
    services.set_attendance_range(manager, start=day, end=day, status=S.PAID_LEAVE, employee=keeper)
    entry = AttendanceEntry.objects.get(employee=keeper, date=day)
    assert (entry.status, entry.overtime_hours) == (S.PAID_LEAVE, Decimal(0))

    def apply(**kwargs: Any) -> Callable[[], Any]:
        values = {"start": first, "end": last, "status": S.ABSENT, **kwargs}
        return lambda: services.set_attendance_range(manager, **values)

    assert _code(apply(start=last, end=first)) == "invalid_range"
    assert _code(apply(end=first + timedelta(days=366))) == "range_too_long"
    assert _code(apply(department="Bakım")) == "no_employees"
    future = timezone.localdate() + timedelta(days=2)
    assert _code(apply(start=future, end=future)) == "future_day"


# --- onay, kilit ve yeniden açma -----------------------------------------------------------------


def test_approve_locks_month_and_writes_every_day(
    manager: User,
    django_capture_on_commit_callbacks: Callable[..., Any],
) -> None:
    received: list[events.Event] = []
    events.subscribe("attendance_month.*")(received.append)
    year, month = _previous_month()
    first, last = selectors.month_bounds(year, month)
    card = employee(manager)
    newcomer = employee(manager, full_name="Can Kaya", hire_date=first + timedelta(days=9))
    leave_day = _workdays(first, last)[0]
    with django_capture_on_commit_callbacks(execute=True):
        services.set_attendance_day(manager, card, leave_day, status=S.ANNUAL_LEAVE)
        record = services.approve_attendance_month(manager, year, month)
    assert record.status == AttendanceMonth.Status.APPROVED
    assert record.approved_by == manager and record.approved_at is not None
    length = (last - first).days + 1
    assert AttendanceEntry.objects.filter(month=record, employee=card).count() == length
    assert AttendanceEntry.objects.filter(month=record, employee=newcomer).count() == length - 9
    assert [event.name for event in received] == [
        "attendance_month.created",
        "attendance_month.approved",
    ]
    assert received[-1].payload["year"] == year and received[-1].payload["month"] == month
    grid = selectors.attendance_grid(year, month)
    assert grid.approved and grid.record == record
    cells = [cell for row in grid.rows for cell in row.cells]
    assert not any(cell.editable for cell in cells)
    assert all(cell.explicit for cell in cells if cell.status)
    assert _cells(grid, card)[leave_day].status == S.ANNUAL_LEAVE

    # Onaylı ayda gün, aralık ve tatil değişmez; ay ikinci kez onaylanmaz.
    assert (
        _code(lambda: services.set_attendance_day(manager, card, first, status=S.ABSENT))
        == "month_locked"
    )
    assert (
        _code(
            lambda: services.set_attendance_range(
                manager, start=first, end=first, status=S.UNPAID_LEAVE
            )
        )
        == "month_locked"
    )
    holiday = {"date": first, "name": "Bayram"}
    assert _code(lambda: services.add_holiday(manager, holiday)) == "month_locked"
    assert _code(lambda: services.approve_attendance_month(manager, year, month)) == "month_locked"

    assert _code(lambda: services.reopen_attendance_month(manager, record, " ")) == (
        "reason_required"
    )
    with django_capture_on_commit_callbacks(execute=True):
        reopened = services.reopen_attendance_month(manager, record, "Rapor eksik girilmiş")
    assert reopened.status == AttendanceMonth.Status.OPEN and reopened.approved_by is None
    assert received[-1].name == "attendance_month.open"
    assert received[-1].payload["reason"] == "Rapor eksik girilmiş"
    note = AuditLog.objects.filter(
        object_id=record.pk, action=AuditLog.Action.STATUS_CHANGE
    ).latest("timestamp")
    assert note.note == "Rapor eksik girilmiş"
    # Günler onaydaki haliyle kalır ve yeniden düzenlenebilir.
    grid = selectors.attendance_grid(year, month)
    assert all(cell.explicit for cell in _cells(grid, card).values())
    services.set_attendance_day(manager, card, first, status=S.SICK_LEAVE)


def test_month_must_be_finished_before_approval(manager: User) -> None:
    today = timezone.localdate()
    code = _code(lambda: services.approve_attendance_month(manager, today.year, today.month))
    assert code == "month_not_finished"


def test_state_machine_transitions(manager: User) -> None:
    assert attendance_month_states.can_transition("OPEN", "APPROVED")
    assert attendance_month_states.can_transition("APPROVED", "OPEN")
    assert _code(lambda: attendance_month_states.assert_can_transition("OPEN", "OPEN")) == (
        "invalid_transition"
    )
    year, month = _previous_month()
    card = employee(manager)
    entry = services.set_attendance_day(manager, card, date(year, month, 3), status=S.WORKED)
    code = _code(lambda: services.reopen_attendance_month(manager, entry.month, "Neden"))
    assert code == "invalid_transition"


def test_fixed_holidays_skip_approved_months(manager: User) -> None:
    year, month = _previous_month()
    services.approve_attendance_month(manager, year, month)
    created = services.add_fixed_holidays(manager, year)
    assert all(holiday.date.month != month for holiday in created)


@pytest.mark.parametrize("role", ["depo", "planlama", "uretim", "laboratuvar", "sevkiyat"])
def test_only_management_and_admin(
    make_user: Callable[..., User], manager: User, role: str
) -> None:
    year, month = _previous_month()
    card = employee(manager)
    other = make_user(role)
    day = date(year, month, 5)
    calls = [
        lambda: services.set_attendance_day(other, card, day, status=S.WORKED),
        lambda: services.set_attendance_range(other, start=day, end=day, status=S.WORKED),
        lambda: services.approve_attendance_month(other, year, month),
        lambda: services.add_holiday(other, {"date": day, "name": "X"}),
        lambda: services.add_fixed_holidays(other, year),
        lambda: services.update_attendance_settings(other, [6]),
    ]
    for call in calls:
        with pytest.raises(PermissionDenied):
            call()
    admin = make_user("admin")
    assert services.set_attendance_day(admin, card, day, status=S.ABSENT).pk


# --- API -----------------------------------------------------------------------------------------


@pytest.fixture
def api(manager: User) -> APIClient:
    client = APIClient()
    client.force_authenticate(manager)
    return client


def test_grid_day_range_approve_and_export_api(api: APIClient, manager: User) -> None:
    year, month = _previous_month()
    first, last = selectors.month_bounds(year, month)
    card = employee(manager, employee_no="S-1")
    url = f"{BASE}/{year}/{month}/"
    body = api.get(url).json()
    assert body["status"] == "OPEN" and body["finished"] is True and body["month_record"] is None
    assert len(body["days"]) == (last - first).days + 1
    row = body["rows"][0]
    assert row["employee"]["full_name"] == "Ayşe Demir"
    assert set(row["cells"][0]) == {
        "date",
        "status",
        "code",
        "overtime_hours",
        "note",
        "explicit",
        "editable",
    }
    assert row["totals"]["missing_days"] == 0

    day = _workdays(first, last)[0].isoformat()
    payload = {"employee": card.pk, "date": day, "status": "WORKED", "overtime_hours": "2.5"}
    response = api.post(f"{BASE}/days/", payload, format="json")
    assert response.status_code == 200, response.json()
    assert response.json()["code"] == "Ç" and response.json()["overtime_hours"] == "2.50"
    payload = {**payload, "status": "ANNUAL_LEAVE", "overtime_hours": "1"}
    response = api.post(f"{BASE}/days/", payload, format="json")
    assert response.status_code == 400 and response.json()["code"] == "overtime_not_allowed"
    week = {
        "employee": card.pk,
        "start": first.isoformat(),
        "end": (first + timedelta(days=6)).isoformat(),
        "status": "SICK_LEAVE",
    }
    response = api.post(f"{BASE}/ranges/", week, format="json")
    assert response.status_code == 200 and response.json() == {"updated": 7}
    body = api.get(url).json()
    assert body["month_record"]["number"].startswith("PU-")
    assert body["rows"][0]["totals"]["days"]["SICK_LEAVE"] == 7

    export = api.get(f"{url}export/")
    assert export.status_code == 200
    assert export["Content-Disposition"] == (
        f'attachment; filename="puantaj-{year}-{month:02d}.xlsx"'
    )
    sheet = load_workbook(BytesIO(export.content))["Puantaj"]
    assert "taslak" in sheet["A1"].value
    assert [sheet["A3"].value, sheet["B3"].value] == ["S-1", "Ayşe Demir"]
    assert sheet.cell(row=3, column=4).value == "R"

    response = api.post(f"{url}approve/")
    assert response.status_code == 200 and response.json()["status"] == "APPROVED"
    assert response.json()["approved_by_name"]
    sheet = load_workbook(BytesIO(api.get(f"{url}export/").content))["Puantaj"]
    assert "onaylı" in sheet["A1"].value
    response = api.post(f"{BASE}/days/", {**payload, "status": "WORKED"}, format="json")
    assert response.status_code == 400 and response.json()["code"] == "month_locked"
    response = api.post(f"{url}reopen/", {"reason": ""}, format="json")
    assert response.status_code == 400 and response.json()["code"] == "validation_error"
    response = api.post(f"{url}reopen/", {"reason": "Düzeltme"}, format="json")
    assert response.status_code == 200 and response.json()["status"] == "OPEN"

    assert api.get(f"{BASE}/{year}/13/").status_code == 404
    assert api.post(f"{BASE}/{year - 1}/{month}/reopen/", {"reason": "X"}).status_code == 404


def test_holiday_and_settings_api(api: APIClient) -> None:
    response = api.post(f"{BASE}/holidays/fixed/", {"year": 2027}, format="json")
    assert response.status_code == 201 and len(response.json()) == 8
    rows = api.get(f"{BASE}/holidays/", {"year": 2027}).json()
    assert [row["date"] for row in rows][:2] == ["2027-01-01", "2027-04-23"]
    holiday = {"date": "2027-03-09", "name": "Ramazan Bayramı 1. gün", "half_day": False}
    response = api.post(f"{BASE}/holidays/", holiday, format="json")
    assert response.status_code == 201, response.json()
    created = response.json()
    response = api.post(f"{BASE}/holidays/", holiday, format="json")
    assert response.status_code == 400 and response.json()["code"] == "holiday_exists"
    response = api.patch(f"{BASE}/holidays/{created['id']}/", {"half_day": True}, format="json")
    assert response.status_code == 200 and response.json()["half_day"] is True
    assert api.post(f"{BASE}/holidays/{created['id']}/remove/").status_code == 200
    assert len(api.get(f"{BASE}/holidays/", {"year": 2027}).json()) == 8
    assert api.delete(f"{BASE}/holidays/{created['id']}/").status_code == 405

    assert api.get(f"{BASE}/settings/").json()["rest_weekdays"] == [6]
    response = api.patch(f"{BASE}/settings/", {"rest_weekdays": [6, 5]}, format="json")
    assert response.status_code == 200 and response.json()["rest_weekdays"] == [5, 6]
    response = api.patch(f"{BASE}/settings/", {"rest_weekdays": []}, format="json")
    assert response.status_code == 400 and response.json()["code"] == "invalid_rest_weekdays"


@pytest.mark.parametrize("role", OTHER_ROLES)
def test_other_roles_cannot_reach_attendance(
    client_for: Callable[..., APIClient], manager: User, role: str
) -> None:
    year, month = _previous_month()
    card = employee(manager)
    client = client_for(role)
    day = {"employee": card.pk, "date": f"{year}-{month:02d}-05", "status": "WORKED"}
    span = {"start": day["date"], "end": day["date"], "status": "WORKED"}
    checks = [
        client.get(f"{BASE}/{year}/{month}/"),
        client.get(f"{BASE}/{year}/{month}/export/"),
        client.post(f"{BASE}/{year}/{month}/approve/"),
        client.post(f"{BASE}/{year}/{month}/reopen/", {"reason": "X"}),
        client.post(f"{BASE}/days/", day, format="json"),
        client.post(f"{BASE}/ranges/", span, format="json"),
        client.get(f"{BASE}/holidays/"),
        client.post(f"{BASE}/holidays/fixed/", {"year": 2027}),
        client.get(f"{BASE}/settings/"),
        client.patch(f"{BASE}/settings/", {"rest_weekdays": [6]}, format="json"),
    ]
    assert [response.status_code for response in checks] == [403] * len(checks)
    assert not AttendanceEntry.objects.exists()
