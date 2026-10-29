"""Puantajın Excel çıktısı (openpyxl): aylık tablo ve toplamlar, günlük fazla mesai, kodlar."""

import io
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.personnel.models import ATTENDANCE_CODES, AttendanceStatus
from apps.personnel.selectors import AttendanceDay, AttendanceGrid

WEEKDAYS = ("Pt", "Sa", "Ça", "Pe", "Cu", "Ct", "Pz")
PERSON_HEADERS = ("Sicil no", "Ad soyad", "Birim")
REST_FILL = PatternFill("solid", fgColor="EEEEEE")
HOLIDAY_FILL = PatternFill("solid", fgColor="FDE2E1")
OVERTIME_HEADERS = (
    (AttendanceStatus.WORKED, "FM iş günü"),
    (AttendanceStatus.WEEKLY_REST, "FM hafta tatili"),
    (AttendanceStatus.PUBLIC_HOLIDAY, "FM resmi tatil"),
)


def _title(grid: AttendanceGrid) -> str:
    if grid.approved and grid.record is not None:
        return f"Puantaj — {grid.label} (onaylı, {grid.record.number})"
    return f"Puantaj — {grid.label} (taslak)"


def _day_header(day: AttendanceDay) -> str:
    return f"{day.date.day} {WEEKDAYS[day.weekday]}"


def _fill(day: AttendanceDay) -> PatternFill | None:
    if day.holiday is not None:
        return HOLIDAY_FILL
    return REST_FILL if day.rest else None


def _sheet(ws: Worksheet, grid: AttendanceGrid, headers: Sequence[str]) -> None:
    """Başlık satırı, sütun genişlikleri ve tatil sütunlarının rengi (veri eklendikten sonra)."""
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    first_day = len(PERSON_HEADERS) + 1
    for index, header in enumerate(headers, start=1):
        width = 5 if first_day <= index < first_day + len(grid.days) else max(10, len(header) + 2)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.column_dimensions["B"].width = 26
    for offset, day in enumerate(grid.days):
        fill = _fill(day)
        if fill is None:
            continue
        for (cell,) in ws.iter_rows(
            min_row=2, min_col=first_day + offset, max_col=first_day + offset
        ):
            cell.fill = fill
    ws.freeze_panes = ws.cell(row=3, column=first_day)


def attendance_xlsx(grid: AttendanceGrid) -> bytes:
    """Aylık puantaj: "Puantaj" (günlük kodlar ve toplamlar), "Fazla mesai" (günlük saat) ve
    "Kodlar" (kod açıklamaları ve ayın tatilleri)."""
    workbook = Workbook()
    ws = workbook.active
    assert ws is not None
    ws.title = "Puantaj"
    statuses = list(AttendanceStatus)
    headers = [
        *PERSON_HEADERS,
        *(_day_header(day) for day in grid.days),
        *(ATTENDANCE_CODES[status] for status in statuses),
        "Eksik gün",
        *(label for _, label in OVERTIME_HEADERS),
        "FM toplam",
    ]
    ws.append([_title(grid)])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(headers)
    for row in grid.rows:
        person = row.employee
        totals = row.totals
        ws.append(
            [
                person.employee_no,
                person.full_name,
                person.department,
                *(cell.code for cell in row.cells),
                *(totals.days[status] for status in statuses),
                totals.missing_days,
                *(totals.overtime[status] for status, _ in OVERTIME_HEADERS),
                totals.overtime_total,
            ]
        )
    _sheet(ws, grid, headers)

    overtime = workbook.create_sheet("Fazla mesai")
    overtime_headers = [*PERSON_HEADERS, *(_day_header(day) for day in grid.days), "Toplam"]
    overtime.append([f"Fazla mesai (saat) — {grid.label}"])
    overtime["A1"].font = Font(bold=True, size=12)
    overtime.append(overtime_headers)
    for row in grid.rows:
        person = row.employee
        overtime.append(
            [
                person.employee_no,
                person.full_name,
                person.department,
                *(cell.overtime_hours or None for cell in row.cells),
                row.totals.overtime_total,
            ]
        )
    _sheet(overtime, grid, overtime_headers)

    legend = workbook.create_sheet("Kodlar")
    legend.append(["Kod", "Durum"])
    for status in statuses:
        legend.append([ATTENDANCE_CODES[status], status.label])
    legend.append([])
    legend.append(["Tarih", "Resmi tatil"])
    for cell in legend[legend.max_row]:
        cell.font = Font(bold=True)
    for day in grid.days:
        if day.holiday is not None:
            suffix = " (yarım gün)" if day.half_day else ""
            legend.append([f"{day.date:%d.%m.%Y}", f"{day.holiday_name}{suffix}"])
    legend.column_dimensions["A"].width = 12
    legend.column_dimensions["B"].width = 40
    for cell in legend[1]:
        cell.font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
