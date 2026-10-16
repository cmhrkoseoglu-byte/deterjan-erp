"""Rapor dosyaları: Excel (openpyxl, salt yazma kipi) ve PDF (reportlab, A4).

Her rapor aynı düzende yazılır: başlık, rapor zamanı, filtreler, özet, tablo. Türkçe karakterler
için analiz sertifikasıyla aynı Vera yazı tipi kullanılır. Tutar sütunları çağıran tarafta
(`services.export_columns`) yetkiye göre süzülmüş gelir.
"""

import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, portrait
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.quality.coa import FONT, FONT_BOLD, _register_fonts
from apps.reporting.reports import Column

NUMBER_FORMATS = {
    "number": "#,##0.####",
    "money": "#,##0.00",
    "percent": "0.00",
    "duration": "0.00",
}
YES, NO = "Evet", "Hayır"


def _local(value: datetime) -> datetime:
    """Excel saat dilimi tutmaz: yerel saate çevrilmiş, saat dilimsiz değer yazılır."""
    return timezone.localtime(value).replace(tzinfo=None)


def number_text(value: Decimal, places: int | None = None) -> str:
    """Türkçe sayı biçimi: binlik nokta, ondalık virgül."""
    text = f"{value:,.{places}f}" if places is not None else f"{value:,.4f}".rstrip("0").rstrip(".")
    return text.replace(",", "_").replace(".", ",").replace("_", ".")


def display(value: Any, kind: str) -> str:
    """PDF ve özet için metin."""
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return YES if value else NO
    if isinstance(value, datetime):
        return f"{timezone.localtime(value):%d.%m.%Y %H:%M}"
    if isinstance(value, date):
        return f"{value:%d.%m.%Y}"
    if isinstance(value, Decimal):
        return number_text(value, 2 if kind in ("money", "percent", "duration") else None)
    if isinstance(value, int) and kind != "text":
        return number_text(Decimal(value), 0)
    return str(value)


def _excel_value(value: Any) -> Any:
    if isinstance(value, bool):
        return YES if value else NO
    if isinstance(value, datetime):
        return _local(value)
    return value


def _header_block(
    sheet: Worksheet,
    title: str,
    filters: Sequence[tuple[str, str]],
    summary: Sequence[dict[str, Any]],
) -> None:
    bold = Font(bold=True)
    title_cell = WriteOnlyCell(sheet, value=title)
    title_cell.font = Font(bold=True, size=14)
    sheet.append([title_cell])
    sheet.append(["Rapor zamanı", f"{timezone.localtime():%d.%m.%Y %H:%M}"])
    for label, text in filters:
        sheet.append([label, text])
    for item in summary:
        label = WriteOnlyCell(sheet, value=item["label"])
        label.font = bold
        sheet.append([label, _excel_value(item["value"])])
    sheet.append([])


def to_xlsx(
    title: str,
    columns: Sequence[Column],
    rows: Iterable[dict[str, Any]],
    *,
    filters: Sequence[tuple[str, str]] = (),
    summary: Sequence[dict[str, Any]] = (),
) -> bytes:
    """Rapor Excel dosyası (tek sayfa). Satırlar akış halinde yazılır (büyük raporlar için)."""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title[:31])
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width + 2
    _header_block(sheet, title, filters, summary)
    bold = Font(bold=True)
    header = []
    for column in columns:
        cell = WriteOnlyCell(sheet, value=column.label)
        cell.font = bold
        header.append(cell)
    sheet.append(header)
    formats = [NUMBER_FORMATS.get(c.kind) for c in columns]
    date_formats = {"date": "DD.MM.YYYY", "datetime": "DD.MM.YYYY HH:MM"}
    for row in rows:
        cells = []
        for column, number_format in zip(columns, formats, strict=True):
            cell = WriteOnlyCell(sheet, value=_excel_value(row.get(column.key)))
            fmt = number_format or date_formats.get(column.kind)
            if fmt:
                cell.number_format = fmt
            cells.append(cell)
        sheet.append(cells)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(
    title: str,
    columns: Sequence[Column],
    rows: Iterable[dict[str, Any]],
    *,
    filters: Sequence[tuple[str, str]] = (),
    summary: Sequence[dict[str, Any]] = (),
    wide: bool = True,
) -> bytes:
    """Rapor PDF'i: A4 (geniş raporlar yatay), başlık satırı her sayfada tekrar eder."""
    _register_fonts()
    pagesize = landscape(A4) if wide else portrait(A4)
    margin = 10 * mm
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=title,
    )
    heading = ParagraphStyle("title", fontName=FONT_BOLD, fontSize=14, leading=18)
    base = ParagraphStyle("base", fontName=FONT, fontSize=8, leading=10)
    cell = ParagraphStyle("cell", fontName=FONT, fontSize=6.5, leading=8)
    head = ParagraphStyle("head", fontName=FONT_BOLD, fontSize=6.5, leading=8)
    story: list[Any] = [
        Paragraph(escape(title), heading),
        Paragraph(escape(f"Rapor zamanı: {timezone.localtime():%d.%m.%Y %H:%M}"), base),
    ]
    for label, text in filters:
        story.append(Paragraph(escape(f"{label}: {text}"), base))
    if summary:
        story.append(Spacer(1, 2 * mm))
        story.append(
            Paragraph(
                escape(
                    " · ".join(f"{s['label']}: {display(s['value'], s['kind'])}" for s in summary)
                ),
                base,
            )
        )
    story.append(Spacer(1, 3 * mm))
    total_width = pagesize[0] - 2 * margin
    weight = sum(c.width for c in columns)
    widths = [total_width * c.width / weight for c in columns]
    data: list[list[Any]] = [[Paragraph(escape(c.label), head) for c in columns]]
    for row in rows:
        data.append([Paragraph(escape(display(row.get(c.key), c.kind)), cell) for c in columns])
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue()
