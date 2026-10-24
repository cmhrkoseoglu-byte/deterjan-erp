"""Raporlama API'si: yetki (rol ve tutar sütunları), sayfalama, Excel/PDF ve arka plan işi."""

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.core.roles import (
    DEPO,
    DOLUM,
    LABORATUVAR,
    PLANLAMA,
    SATIN_ALMA,
    SEVKIYAT,
    URETIM,
    YONETIM,
)
from apps.inventory.services import set_lot_cost
from apps.inventory.tests.factories import stocked_lot
from apps.reporting import services
from apps.reporting.models import ReportJob
from apps.workflow.models import Notification

pytestmark = pytest.mark.django_db
BASE = "/api/v1/reporting"

DASHBOARD_ACCESS = {
    "management": ({YONETIM, PLANLAMA}, DEPO),
    "warehouse": ({DEPO, PLANLAMA, YONETIM}, DOLUM),
    "lab": ({LABORATUVAR, PLANLAMA, YONETIM}, SEVKIYAT),
    "production": ({URETIM, PLANLAMA, YONETIM}, LABORATUVAR),
    "shipping": ({SEVKIYAT, DEPO, PLANLAMA, YONETIM}, URETIM),
}
REPORT_ACCESS = {
    "stock-status": (DEPO, DOLUM),
    "stock-movements": (SATIN_ALMA, SEVKIYAT),
    "production": (URETIM, DEPO),
    "filling": (DOLUM, URETIM),
    "quality": (LABORATUVAR, DEPO),
    "shipments": (SEVKIYAT, DEPO),
    "task-durations": (PLANLAMA, DEPO),
}


def costed_lot(quantity: str, unit_cost: str):
    lot = stocked_lot(quantity)
    return set_lot_cost(None, lot, Decimal(unit_cost), estimated=False)


def test_anonymous_gets_401(api_client):
    assert api_client.get(f"{BASE}/dashboards/management/").status_code == 401
    assert api_client.get(f"{BASE}/reports/").status_code == 401


@pytest.mark.parametrize("code", DASHBOARD_ACCESS)
def test_dashboard_roles(client_for, code):
    allowed, denied = DASHBOARD_ACCESS[code]
    for role in allowed:
        assert client_for(role).get(f"{BASE}/dashboards/{code}/").status_code == 200, role
    response = client_for(denied).get(f"{BASE}/dashboards/{code}/")
    assert response.status_code == 403
    assert client_for("admin").get(f"{BASE}/dashboards/{code}/").status_code == 200


def test_dashboard_list_and_unknown(client_for):
    codes = [d["code"] for d in client_for(DEPO).get(f"{BASE}/dashboards/").json()]
    assert codes == ["warehouse", "shipping"]
    assert client_for(YONETIM).get(f"{BASE}/dashboards/nope/").status_code == 404


def test_stock_value_only_for_management(client_for):
    costed_lot("10", "2.5")
    planning = client_for(PLANLAMA).get(f"{BASE}/dashboards/management/").json()
    assert planning["stock_value"] is None
    assert planning["approvals"]["purchase_orders"] == {"count": 0, "total_try": None}
    data = client_for(YONETIM).get(f"{BASE}/dashboards/management/").json()
    assert data["stock_value"]["total"] == "25.00"
    assert data["approvals"]["purchase_orders"] == {"count": 0, "total_try": "0"}
    assert len(data["daily"]) == 14


@pytest.mark.parametrize("code", REPORT_ACCESS)
def test_report_roles(client_for, code):
    allowed, denied = REPORT_ACCESS[code]
    assert client_for(allowed).get(f"{BASE}/reports/{code}/").status_code == 200
    assert client_for(denied).get(f"{BASE}/reports/{code}/").status_code == 403
    response = client_for(denied).post(
        f"{BASE}/reports/{code}/export/", {"file_format": "xlsx"}, format="json"
    )
    assert response.status_code == 403


def test_report_list_is_narrowed_by_role(client_for):
    codes = {r["code"] for r in client_for(LABORATUVAR).get(f"{BASE}/reports/").json()}
    assert codes == {"quality"}
    assert len(client_for("admin").get(f"{BASE}/reports/").json()) == 7


def test_cost_columns_hidden_from_depo(client_for):
    lot = costed_lot("10", "2.5")
    depo = client_for(DEPO).get(f"{BASE}/reports/stock-status/").json()
    assert "value" not in {c["key"] for c in depo["columns"]}
    assert "value" not in depo["results"][0] and "unit_cost" not in depo["results"][0]
    assert all(s["kind"] != "money" for s in depo["summary"])
    manager = client_for(YONETIM).get(f"{BASE}/reports/stock-status/").json()
    [row] = [r for r in manager["results"] if r["lot_no"] == lot.lot_no]
    assert row["value"] == "25.00" and row["unit_cost"] == "2.5000"
    # Excel'de de tutar sütunu yok.
    content = client_for(DEPO).post(
        f"{BASE}/reports/stock-status/export/", {"file_format": "xlsx"}, format="json"
    )
    sheet = load_workbook(io.BytesIO(content.content)).active
    values = [c for row in sheet.iter_rows(values_only=True) for c in row if c]
    assert "Stok değeri (TL)" not in values and "Kullanılabilir" in values


def test_filters_and_paging(client_for):
    lot = stocked_lot("10")
    stocked_lot("5")
    client = client_for(DEPO)
    page = client.get(f"{BASE}/reports/stock-status/", {"page_size": 1}).json()
    assert page["count"] == 2 and len(page["results"]) == 1
    only = client.get(f"{BASE}/reports/stock-status/", {"material": lot.material_id}).json()
    assert [r["lot_no"] for r in only["results"]] == [lot.lot_no]
    bad = client.get(
        f"{BASE}/reports/stock-movements/", {"date_from": "2026-05-02", "date_to": "2026-05-01"}
    )
    assert bad.status_code == 400
    movements = client.get(f"{BASE}/reports/stock-movements/").json()
    assert movements["filters"]["date_from"]  # varsayılan son 30 gün
    assert movements["count"] == 2


def test_sync_export_xlsx_and_pdf(client_for):
    stocked_lot("10")
    client = client_for(DEPO)
    xlsx = client.post(
        f"{BASE}/reports/stock-movements/export/", {"file_format": "xlsx"}, format="json"
    )
    assert xlsx.status_code == 200
    assert xlsx["Content-Disposition"].startswith('attachment; filename="stock-movements-')
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet["A1"].value == "Stok hareketleri"
    pdf = client.post(
        f"{BASE}/reports/stock-movements/export/", {"file_format": "pdf"}, format="json"
    )
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    bad = client.post(f"{BASE}/reports/stock-movements/export/", {"file_format": "csv"})
    assert bad.status_code == 400


def test_large_export_runs_in_background(
    client_for, make_user, monkeypatch, django_capture_on_commit_callbacks
):
    monkeypatch.setattr(services, "SYNC_ROW_LIMIT", 1)
    stocked_lot("10")
    stocked_lot("5")
    user = make_user(DEPO)
    client = client_for(DEPO)
    client.force_authenticate(user)
    with django_capture_on_commit_callbacks(execute=True):
        response = client.post(
            f"{BASE}/reports/stock-status/export/", {"file_format": "xlsx"}, format="json"
        )
    assert response.status_code == 202
    job = ReportJob.objects.get(pk=response.json()["id"])
    assert job.status == ReportJob.Status.DONE and job.row_count == 2
    assert job.created_by == user
    note = Notification.objects.get(recipient=user, rule_code="report_ready")
    assert note.title == "Rapor hazır: Stok durumu"
    assert note.related_object_id == job.pk

    listed = client.get(f"{BASE}/jobs/").json()["results"]
    assert [j["id"] for j in listed] == [job.pk] and listed[0]["can_download"]
    download = client.get(f"{BASE}/jobs/{job.pk}/download/")
    assert download.status_code == 200
    sheet = load_workbook(io.BytesIO(b"".join(download.streaming_content))).active
    assert sheet["A1"].value == "Stok durumu"
    # Başkası göremez ve indiremez.
    other = client_for(DEPO)
    assert other.get(f"{BASE}/jobs/{job.pk}/").status_code == 404
    assert other.get(f"{BASE}/jobs/{job.pk}/download/").status_code == 404


def test_pdf_row_limit(client_for, monkeypatch):
    monkeypatch.setattr(services, "PDF_ROW_LIMIT", 1)
    stocked_lot("10")
    stocked_lot("5")
    response = client_for(DEPO).post(
        f"{BASE}/reports/stock-status/export/", {"file_format": "pdf"}, format="json"
    )
    assert response.status_code == 400
    assert response.json()["code"] == "pdf_too_large"
