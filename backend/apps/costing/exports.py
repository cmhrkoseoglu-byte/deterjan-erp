"""Maliyet raporlarının Excel çıktısı (openpyxl). Her rapor tek sayfa; stok değeri iki sayfa
(malzeme özeti ve lot dökümü), personel maliyeti üç sayfa (kişiler, aşamalar, özet)."""

import io
from collections.abc import Sequence
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.costing.selectors import PersonnelCostReport

TYPE_LABELS = {"RAW": "Hammadde", "PACKAGING": "Ambalaj", "BULK": "Dökme", "FINISHED": "Mamul"}
STAGE_LABELS = {"PRODUCTION": "Üretim (bulk)", "FILLING": "Dolum (mamul)", "": "Genel gider"}

Column = tuple[str, str]  # (alan, başlık)

COST_TREND: Sequence[Column] = (
    ("material_code", "Ürün kodu"),
    ("material_name", "Ürün"),
    ("month", "Ay"),
    ("lot_count", "Lot sayısı"),
    ("quantity", "Miktar"),
    ("unit", "Birim"),
    ("total_cost", "Toplam maliyet (TL)"),
    ("unit_cost", "Birim maliyet (TL)"),
    ("has_estimated", "Tahmini içerir"),
)
BATCH_VARIANCE: Sequence[Column] = (
    ("lot_no", "Şarj / lot"),
    ("order_number", "Üretim emri"),
    ("product_code", "Ürün kodu"),
    ("product_name", "Ürün"),
    ("cost_date", "Tarih"),
    ("planned_quantity", "Şarj miktarı"),
    ("actual_quantity", "Gerçek çıktı"),
    ("unit", "Birim"),
    ("theoretical_material_cost", "Teorik malzeme (TL)"),
    ("actual_material_cost", "Fiili malzeme (TL)"),
    ("overhead_cost", "Genel gider (TL)"),
    ("total_cost", "Toplam (TL)"),
    ("theoretical_unit_cost", "Teorik birim (TL)"),
    ("actual_material_unit_cost", "Fiili birim malzeme (TL)"),
    ("unit_cost", "Fiili birim toplam (TL)"),
    ("variance_pct", "Fark (%)"),
    ("is_estimated", "Tahmini"),
)
SCRAP_COST: Sequence[Column] = (
    ("lot_no", "Dolum lotu"),
    ("product_code", "Ürün kodu"),
    ("product_name", "Ürün"),
    ("cost_date", "Tarih"),
    ("good_units", "Sağlam adet"),
    ("material_scrap_cost", "Bulk ve ambalaj firesi (TL)"),
    ("finished_scrap_cost", "Mamul firesi payı (TL)"),
    ("total_scrap_cost", "Toplam fire (TL)"),
    ("total_cost", "Lot maliyeti (TL)"),
    ("scrap_pct", "Fire oranı (%)"),
)
TOP_MATERIALS: Sequence[Column] = (
    ("material_code", "Malzeme kodu"),
    ("material_name", "Malzeme"),
    ("material_type", "Tür"),
    ("quantity", "Tüketim"),
    ("unit", "Birim"),
    ("amount", "Tutar (TL)"),
    ("average_unit_cost", "Ortalama birim (TL)"),
    ("share_pct", "Pay (%)"),
    ("has_missing", "Eksik maliyet"),
)
STOCK_VALUE: Sequence[Column] = (
    ("material_code", "Malzeme kodu"),
    ("material_name", "Malzeme"),
    ("material_type", "Tür"),
    ("quantity", "Miktar"),
    ("unit", "Birim"),
    ("value", "Değer (TL)"),
    ("average_unit_cost", "Ortalama birim (TL)"),
    ("uncosted_quantity", "Maliyetsiz miktar"),
    ("estimated_quantity", "Tahmini maliyetli miktar"),
)
STOCK_VALUE_LOTS: Sequence[Column] = (
    ("material_code", "Malzeme kodu"),
    ("material_name", "Malzeme"),
    ("lot_no", "Lot no"),
    ("warehouse", "Depo"),
    ("location", "Lokasyon"),
    ("quantity", "Miktar"),
    ("unit", "Birim"),
    ("unit_cost", "Birim maliyet (TL)"),
    ("value", "Değer (TL)"),
    ("is_estimated", "Tahmini"),
)


def _cell(name: str, value: Any) -> Any:
    if isinstance(value, bool):
        return "Evet" if value else ""
    if name == "material_type":
        return TYPE_LABELS.get(value, value)
    if name == "stage":
        return STAGE_LABELS.get(value, value)
    if name == "month" and isinstance(value, date):
        return f"{value:%m.%Y}"
    return value


def _sheet(ws: Worksheet, title: str, columns: Sequence[Column], rows: list[dict]) -> None:
    ws.title = title[:31]
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_cell(name, row.get(name)) for name, _ in columns])
    for index, (_, label) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = max(12, len(label) + 2)
    ws.freeze_panes = "A2"


def to_xlsx(sheets: Sequence[tuple[str, Sequence[Column], list[dict]]]) -> bytes:
    workbook = Workbook()
    for index, (title, columns, rows) in enumerate(sheets):
        ws = workbook.active if index == 0 else workbook.create_sheet()
        assert ws is not None
        _sheet(ws, title, columns, rows)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


PERSONNEL_COST: Sequence[Column] = (
    ("employee_no", "Sicil no"),
    ("full_name", "Ad soyad"),
    ("department", "Birim"),
    ("stage", "Maliyet merkezi"),
    ("gross_monthly", "Brüt aylık ücret (TL)"),
    ("employed_days", "İşte olduğu gün"),
    ("unpaid_days", "Ücretsiz gün"),
    ("paid_days", "Ödenen gün"),
    ("overtime_hours", "Fazla mesai (saat)"),
    ("gross_salary", "Brüt ücret (TL)"),
    ("overtime_pay", "Fazla mesai (TL)"),
    ("employer_share", "İşveren payı (TL)"),
    ("total_cost", "Toplam maliyet (TL)"),
    ("missing_wage", "Ücret eksik"),
)
PERSONNEL_STAGES: Sequence[Column] = (
    ("stage", "Maliyet merkezi"),
    ("employee_count", "Kişi"),
    ("cost", "Personel maliyeti (TL)"),
    ("kg", "Üretilen (kg)"),
    ("unconverted", "Kg'a çevrilemeyen malzeme"),
    ("suggested_rate", "Gerçekleşen işçilik (TL/kg)"),
    ("current_rate", "Geçerli işçilik oranı (TL/kg)"),
    ("missing_wage_count", "Ücreti eksik kişi"),
)
SUMMARY: Sequence[Column] = (("label", "Bilgi"), ("value", "Değer"))


def personnel_cost_xlsx(report: PersonnelCostReport) -> bytes:
    """Aylık personel maliyeti: kişiler, aşamalar (işçilik önerisiyle) ve özet."""
    if report.attendance_approved:
        attendance = f"Onaylı ({report.attendance_number})"
    else:
        attendance = "Açık (taslak)"
    summary: list[dict[str, Any]] = [
        {"label": "Dönem", "value": report.label},
        {"label": "Puantaj", "value": attendance},
        {"label": "İşveren payı (%)", "value": report.employer_rate},
        {"label": "Brüt ücret toplamı (TL)", "value": report.gross_salary},
        {"label": "Fazla mesai toplamı (TL)", "value": report.overtime_pay},
        {"label": "İşveren payı toplamı (TL)", "value": report.employer_share},
        {"label": "Toplam personel maliyeti (TL)", "value": report.total_cost},
        {"label": "Ücreti eksik kişi", "value": report.missing_wage_count},
    ]
    rows = [
        {
            "employee_no": row.employee.employee_no,
            "full_name": row.employee.full_name,
            "department": row.employee.department,
            "stage": row.stage,
            "gross_monthly": row.gross_monthly,
            "employed_days": row.employed_days,
            "unpaid_days": row.unpaid_days,
            "paid_days": row.paid_days,
            "overtime_hours": row.overtime_hours,
            "gross_salary": row.gross_salary,
            "overtime_pay": row.overtime_pay,
            "employer_share": row.employer_share,
            "total_cost": row.total_cost,
            "missing_wage": row.missing_wage,
        }
        for row in report.rows
    ]
    stages = [
        {
            "stage": item.stage,
            "employee_count": item.employee_count,
            "cost": item.cost,
            "kg": item.kg,
            "unconverted": item.unconverted,
            "suggested_rate": item.suggested_rate,
            "current_rate": item.current_rate,
            "missing_wage_count": item.missing_wage_count,
        }
        for item in report.stages
    ]
    return to_xlsx(
        [
            ("Personel maliyeti", PERSONNEL_COST, rows),
            ("Aşamalar", PERSONNEL_STAGES, stages),
            ("Özet", SUMMARY, summary),
        ]
    )
