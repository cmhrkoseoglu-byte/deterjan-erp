"""Personel maliyeti: ücretler, 30 gün usulü aylık hesap, fazla mesai, aşama toplamları, üretilen kg
ve işçilik oranı önerisi (docs/DECISIONS.md "Personel maliyeti")."""

from collections.abc import Callable
from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from functools import partial
from io import BytesIO
from typing import Any

import pytest
from django.core.exceptions import PermissionDenied
from django.db import IntegrityError
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core.exceptions import BusinessRuleError
from apps.core.models import AuditLog, User
from apps.core.roles import ADMIN, ALL_ROLES, YONETIM
from apps.costing import selectors, services
from apps.costing.api.serializers import WageSummarySerializer
from apps.costing.models import (
    CostRate,
    CostStage,
    DepartmentCostCenter,
    EmployeeWage,
    PersonnelCostSettings,
)
from apps.costing.tests.factories import build_chain
from apps.filling.models import FillingRun
from apps.personnel import services as personnel
from apps.personnel.models import AttendanceStatus, Employee
from apps.personnel.tests.factories import employee
from apps.production.models import ProductionBatch

D = Decimal
S = AttendanceStatus
URL = "/api/v1/costing"
# Yönetim ve admin dışındaki bütün roller (dolum ve sevkiyat dahil).
OTHER_ROLES = tuple(role for role in ALL_ROLES if role not in (ADMIN, YONETIM))

pytestmark = pytest.mark.usefixtures("invariants")


@pytest.fixture
def manager(staff: dict[str, User]) -> User:
    return staff[YONETIM]


def _wage(actor: User, person: Employee, amount: str, valid_from: date) -> EmployeeWage:
    return services.add_wage(
        actor, {"employee": person, "gross_monthly": D(amount), "valid_from": valid_from}
    )


def _row(year: int, month: int, person: Employee) -> selectors.PersonnelCostRow:
    report = selectors.personnel_cost_report(year, month)
    return next(row for row in report.rows if row.employee.pk == person.pk)


def _code(call: Callable[[], Any]) -> str:
    with pytest.raises(BusinessRuleError) as exc:
        call()
    return exc.value.code


def _sheet_rows(sheet: Any) -> list[dict[str, Any]]:
    """Excel sayfasını başlık → değer sözlüklerine çevirir."""
    rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


# --- aylık hesap -----------------------------------------------------------------------------


def test_full_month_and_unpaid_days(manager: User) -> None:
    """Ağustos 2026 (31 gün): tam ay 30 gün sayılır; ücretsiz izin, devamsızlık ve rapor günlük
    ücretle (30.000 / 30 = 1.000) düşülür; mazeret ve yıllık izin ücretlidir."""
    person = employee(manager)
    _wage(manager, person, "30000", date(2020, 1, 1))
    row = _row(2026, 8, person)
    assert (row.employed_days, row.paid_days, row.unpaid_days) == (31, 30, 0)
    assert row.gross_salary == D("30000.00") and row.overtime_pay == D("0.00")
    assert row.employer_share == D("5325.00") and row.total_cost == D("35325.00")
    assert row.gross_monthly == D("30000.00") and not row.missing_wage

    for day, status in (
        (3, S.UNPAID_LEAVE),
        (4, S.ABSENT),
        (5, S.SICK_LEAVE),
        (6, S.PAID_LEAVE),
        (7, S.ANNUAL_LEAVE),
    ):
        personnel.set_attendance_day(manager, person, date(2026, 8, day), status=status)
    row = _row(2026, 8, person)
    assert (row.paid_days, row.unpaid_days) == (27, 3)
    assert row.gross_salary == D("27000.00")
    assert row.employer_share == D("4792.50") and row.total_cost == D("31792.50")


def test_february_and_partial_months(manager: User) -> None:
    """Şubat 2026 (28 gün): tam ay yine aylık brüt; ay içinde giren ya da ayrılan takvim günü kadar
    ücret alır."""
    whole = employee(manager)
    newcomer = employee(manager, full_name="Can Kaya", hire_date=date(2026, 2, 15))
    leaver = employee(manager, full_name="Ece Ak", hire_date=date(2020, 1, 1))
    personnel.terminate_employee(manager, leaver, date(2026, 2, 10))
    for person in (whole, newcomer, leaver):
        _wage(manager, person, "30000", date(2020, 1, 1))
    assert _row(2026, 2, whole).gross_salary == D("30000.00")
    assert _row(2026, 2, whole).paid_days == 30
    assert (_row(2026, 2, newcomer).paid_days, _row(2026, 2, newcomer).gross_salary) == (
        14,
        D("14000.00"),
    )
    assert (_row(2026, 2, leaver).paid_days, _row(2026, 2, leaver).gross_salary) == (
        10,
        D("10000.00"),
    )
    # Ayın tamamı ücretsiz izin: 30 - 28 = 2 gün kalmaz, ücret sıfırdır.
    personnel.set_attendance_range(
        manager,
        start=date(2026, 2, 1),
        end=date(2026, 2, 28),
        status=S.UNPAID_LEAVE,
        employee=whole,
    )
    row = _row(2026, 2, whole)
    assert (row.paid_days, row.unpaid_days, row.gross_salary) == (0, 28, D("0.00"))
    assert row.total_cost == D("0.00")


def test_thirty_unpaid_days_in_a_31_day_month_pay_the_worked_day(manager: User) -> None:
    """Ağustos 2026 (31 gün): 1'inde çalışıp 2–31 ücretsiz izinli olan 30 - 30 = 0 değil, ücretli
    takvim günü kadar (1 gün) alır; 31 günün tamamı ücretsizse sıfır."""
    worker = employee(manager)
    idle = employee(manager, full_name="Can Kaya")
    for person in (worker, idle):
        _wage(manager, person, "30000", date(2020, 1, 1))
    personnel.set_attendance_range(
        manager,
        start=date(2026, 8, 2),
        end=date(2026, 8, 31),
        status=S.UNPAID_LEAVE,
        employee=worker,
    )
    personnel.set_attendance_range(
        manager, start=date(2026, 8, 1), end=date(2026, 8, 31), status=S.UNPAID_LEAVE, employee=idle
    )
    row = _row(2026, 8, worker)
    assert (row.paid_days, row.unpaid_days, row.gross_salary) == (1, 30, D("1000.00"))
    row = _row(2026, 8, idle)
    assert (row.paid_days, row.gross_salary) == (0, D("0.00"))


def test_half_kurus_rounds_up_without_drift(manager: User) -> None:
    """Eylül 2026: 147.404,15 TL'den 22 Eylül'de 161.300,50 TL'ye zam. Kesin değer
    (21 × 147.404,15 + 9 × 161.300,50) / 30 = 151.573,055 → yukarı yuvarlanır: 151.573,06."""
    person = employee(manager)
    _wage(manager, person, "147404.15", date(2020, 1, 1))
    _wage(manager, person, "161300.50", date(2026, 9, 22))
    assert _row(2026, 9, person).gross_salary == D("151573.06")


def test_overtime_by_day_type(manager: User) -> None:
    """Eylül 2026, brüt 22.500 → saat ücreti 100. İş günü ve hafta tatilinde 1,5 kat, resmi tatilde
    çalışılan saate ayrıca 1 kat."""
    person = employee(manager)
    _wage(manager, person, "22500", date(2020, 1, 1))
    personnel.add_holiday(manager, {"date": date(2026, 9, 15), "name": "Bayram"})
    personnel.set_attendance_day(
        manager, person, date(2026, 9, 1), status=S.WORKED, overtime_hours=D(2)
    )
    personnel.set_attendance_day(
        manager, person, date(2026, 9, 6), status=S.WEEKLY_REST, overtime_hours=D(4)
    )
    personnel.set_attendance_day(
        manager, person, date(2026, 9, 15), status=S.PUBLIC_HOLIDAY, overtime_hours=D(3)
    )
    row = _row(2026, 9, person)
    assert row.overtime_hours == D(9)
    assert row.overtime_pay == D("1200.00")  # 300 + 600 + 300
    assert row.gross_salary == D("22500.00") and row.gross_total == D("23700.00")
    assert row.employer_share == D("4206.75") and row.total_cost == D("27906.75")

    services.update_personnel_cost_settings(manager, D("22.75"))
    row = _row(2026, 9, person)
    assert row.employer_share == D("5391.75") and row.total_cost == D("29091.75")


def test_mid_month_raise_and_missing_wage(manager: User) -> None:
    """Eylül 2026 (30 gün): 16 Eylül'den geçerli zamda her gün kendi ücretiyle; ücreti 11 Eylül'de
    başlayan kişinin ilk 10 günü sıfır sayılır ve kişi "ücret eksik" işaretlenir."""
    raised = employee(manager)
    _wage(manager, raised, "30000", date(2020, 1, 1))
    _wage(manager, raised, "36000", date(2026, 9, 16))
    late = employee(manager, full_name="Can Kaya")
    _wage(manager, late, "30000", date(2026, 9, 11))
    nobody = employee(manager, full_name="Ece Ak")

    row = _row(2026, 9, raised)
    assert row.gross_salary == D("33000.00") and row.gross_monthly == D("36000.00")
    assert not row.missing_wage
    row = _row(2026, 9, late)
    assert row.gross_salary == D("20000.00") and row.missing_wage
    row = _row(2026, 9, nobody)
    assert row.total_cost == D("0.00") and row.gross_monthly is None and row.missing_wage
    assert selectors.personnel_cost_report(2026, 9).missing_wage_count == 2


def test_stage_totals_follow_department_mapping(manager: User) -> None:
    filler = employee(manager, department="Dolum")
    maker = employee(manager, full_name="Can Kaya", department="Üretim")
    keeper = employee(manager, full_name="Ece Ak", department="Depo")
    for person, amount in ((filler, "30000"), (maker, "40000"), (keeper, "20000")):
        _wage(manager, person, amount, date(2020, 1, 1))
    services.set_department_cost_center(manager, " dolum ", CostStage.FILLING)
    services.set_department_cost_center(manager, "Üretim", CostStage.PRODUCTION)
    CostRate.objects.create(
        stage=CostStage.PRODUCTION,
        kind=CostRate.Kind.LABOR,
        rate_per_kg=D("2.5"),
        valid_from=date(2026, 9, 1),
    )

    report = selectors.personnel_cost_report(2026, 9)
    production = report.stage_cost(CostStage.PRODUCTION)
    filling = report.stage_cost(CostStage.FILLING)
    general = report.stage_cost(selectors.GENERAL)
    assert (production.cost, production.employee_count) == (D("47100.00"), 1)
    assert (filling.cost, filling.employee_count) == (D("35325.00"), 1)
    assert (general.cost, general.employee_count) == (D("23550.00"), 1)
    assert report.total_cost == D("105975.00")
    assert production.current_rate == D("2.5") and filling.current_rate is None
    # Bu ay üretim yok: öneri hesaplanmaz.
    assert production.kg == 0 and production.suggested_rate is None
    assert general.kg is None and general.suggested_rate is None
    # Eşleştirme büyük/küçük harf farkıyla güncellenir (yeni kayıt açılmaz); boş aşama genel
    # giderdir.
    services.set_department_cost_center(manager, "DOLUM", "")
    assert DepartmentCostCenter.objects.count() == 2
    assert DepartmentCostCenter.objects.get(key="dolum").stage == ""
    rows = {row["department"]: row["stage"] for row in selectors.cost_centers()}
    assert rows == {"Depo": "", "Dolum": "", "Üretim": "PRODUCTION"}
    report = selectors.personnel_cost_report(2026, 9)
    assert report.stage_cost(selectors.GENERAL).employee_count == 2
    # Türkçe İ/ı: "İDARİ İŞLER" ile "İdari İşler" aynı birimdir.
    clerk = employee(manager, full_name="Deniz Su", department="İdari İşler")
    _wage(manager, clerk, "25000", date(2020, 1, 1))
    services.set_department_cost_center(manager, "İDARİ İŞLER", CostStage.PRODUCTION)
    services.set_department_cost_center(manager, "idari işler", CostStage.PRODUCTION)
    assert DepartmentCostCenter.objects.count() == 3
    assert _row(2026, 9, clerk).stage == CostStage.PRODUCTION


# --- üretilen kg ve işçilik oranı ------------------------------------------------------------


def test_produced_kg_and_labor_rate_from_approved_month(
    staff: dict[str, User], commit: Callable[[], Any], manager: User
) -> None:
    chain = build_chain(staff, commit)
    assert chain.batch is not None and chain.run is not None
    today = timezone.localdate()
    previous = today.replace(day=1) - timedelta(days=1)
    moment = timezone.make_aware(datetime.combine(previous.replace(day=15), time(12)))
    ProductionBatch.objects.filter(pk=chain.batch.pk).update(finished_at=moment)
    FillingRun.objects.filter(pk=chain.run.pk).update(finished_at=moment)

    produced = selectors.monthly_produced_kg(previous.year, previous.month)
    assert produced[CostStage.PRODUCTION].kg == D("980")
    assert produced[CostStage.FILLING].kg == D("780")
    assert produced[CostStage.PRODUCTION].unconverted == 0
    assert selectors.monthly_produced_kg(today.year, today.month)[CostStage.PRODUCTION].kg == 0

    maker = employee(manager, full_name="Can Kaya", department="Üretim")
    _wage(manager, maker, "45000", date(2020, 1, 1))
    services.set_department_cost_center(manager, "Üretim", CostStage.PRODUCTION)
    report = selectors.personnel_cost_report(previous.year, previous.month)
    production = report.stage_cost(CostStage.PRODUCTION)
    assert production.cost == D("52987.50") and production.kg == D("980")
    assert production.suggested_rate == D("54.0689")  # 52.987,50 / 980

    start = today + timedelta(days=2)

    def add(stage: str = CostStage.PRODUCTION, **when: int) -> Callable[[], Any]:
        year = when.get("year", previous.year)
        month = when.get("month", previous.month)
        return lambda: services.add_labor_rate_from_personnel_cost(
            manager, year, month, stage, start
        )

    assert _code(add()) == "attendance_not_approved"
    personnel.approve_attendance_month(manager, previous.year, previous.month)
    rate = add()()
    assert (rate.stage, rate.kind, rate.rate_per_kg, rate.valid_from) == (
        CostStage.PRODUCTION,
        CostRate.Kind.LABOR,
        D("54.0689"),
        start,
    )
    assert rate.note == f"Personel maliyeti {report.label}: 52.987,50 TL / 980,00 kg"
    assert rate.created_by == manager
    assert _code(add()) == "duplicate_rate"
    assert _code(add(CostStage.FILLING)) == "no_labor_rate"  # dolumda personel yok
    assert _code(add("")) == "invalid_stage"
    assert _code(add(year=today.year, month=today.month)) == "attendance_not_approved"
    # Personel maliyeti var ama o ay üretim yok: oran hesaplanamaz.
    earlier = previous.replace(day=1) - timedelta(days=1)
    personnel.approve_attendance_month(manager, earlier.year, earlier.month)
    assert _code(add(year=earlier.year, month=earlier.month)) == "no_labor_rate"
    # API: 201 ve oran kaydı.
    client = APIClient()
    client.force_authenticate(manager)
    response = client.post(
        f"{URL}/personnel-costs/{previous.year}/{previous.month}/labor-rate/",
        {"stage": "PRODUCTION", "valid_from": str(start + timedelta(days=1))},
        format="json",
    )
    assert response.status_code == 201, response.json()
    assert response.json()["rate_per_kg"] == "54.0689" and response.json()["kind"] == "LABOR"


def test_produced_kg_follows_istanbul_month_and_completed_status(
    staff: dict[str, User], commit: Callable[[], Any]
) -> None:
    chain = build_chain(staff, commit)
    assert chain.batch is not None and chain.run is not None
    # 30.09.2026 21:30 UTC = 01.10.2026 00:30 İstanbul; 20:30 UTC = 30.09.2026 23:30 İstanbul.
    ProductionBatch.objects.filter(pk=chain.batch.pk).update(
        finished_at=datetime(2026, 9, 30, 21, 30, tzinfo=UTC)
    )
    FillingRun.objects.filter(pk=chain.run.pk).update(
        finished_at=datetime(2026, 9, 30, 20, 30, tzinfo=UTC)
    )
    september = selectors.monthly_produced_kg(2026, 9)
    october = selectors.monthly_produced_kg(2026, 10)
    assert (september[CostStage.PRODUCTION].kg, october[CostStage.PRODUCTION].kg) == (0, D("980"))
    assert (september[CostStage.FILLING].kg, october[CostStage.FILLING].kg) == (D("780"), 0)
    # Bitmemiş şarj ve parti sayılmaz.
    ProductionBatch.objects.filter(pk=chain.batch.pk).update(
        status=ProductionBatch.Status.IN_PROGRESS
    )
    FillingRun.objects.filter(pk=chain.run.pk).update(status=FillingRun.Status.IN_PROGRESS)
    assert selectors.monthly_produced_kg(2026, 10)[CostStage.PRODUCTION].kg == 0
    assert selectors.monthly_produced_kg(2026, 9)[CostStage.FILLING].kg == 0


def test_approved_month_uses_stored_days_within_employment(manager: User) -> None:
    """Onaylı ay kayıtlı günlerle hesaplanır: onaydan sonra açılan kart o aya girmez. Sonradan
    girilen çıkış tarihinden sonraki kayıtlı günler sayılmaz (işte olduğu takvim günü kadar)."""
    person = employee(manager)
    _wage(manager, person, "30000", date(2020, 1, 1))
    personnel.set_attendance_day(manager, person, date(2026, 9, 3), status=S.UNPAID_LEAVE)
    personnel.approve_attendance_month(manager, 2026, 9)
    newcomer = employee(manager, full_name="Can Kaya")
    report = selectors.personnel_cost_report(2026, 9)
    assert [row.employee.pk for row in report.rows] == [person.pk]
    assert report.rows[0].gross_salary == D("29000.00")
    personnel.terminate_employee(manager, person, date(2026, 9, 20))
    row = _row(2026, 9, person)
    assert (row.employed_days, row.unpaid_days, row.paid_days) == (20, 1, 19)
    assert row.gross_salary == D("19000.00")
    october = selectors.personnel_cost_report(2026, 10)
    assert [row.employee.pk for row in october.rows] == [newcomer.pk]


# --- ücretler, ayarlar, yetki ----------------------------------------------------------------


def test_wage_rules(manager: User) -> None:
    person = employee(manager)
    wage = _wage(manager, person, "30000", date(2026, 1, 1))
    assert wage.created_by == manager and wage.is_active

    def add(**values: Any) -> Callable[[], Any]:
        data = {"employee": person, "gross_monthly": D(1), "valid_from": date(2026, 1, 1)}
        return lambda: services.add_wage(manager, {**data, **values})

    assert _code(add()) == "duplicate_wage"
    assert _code(add(gross_monthly=D(0), valid_from=date(2026, 2, 1))) == "invalid_wage"
    assert _code(add(gross_monthly=D("100000000.01"), valid_from=date(2026, 2, 1))) == (
        "invalid_wage"
    )
    assert _code(add(valid_from=date(1999, 12, 31))) == "invalid_valid_from"
    assert _code(add(valid_from=date(2101, 1, 1))) == "invalid_valid_from"
    assert _code(add(valid_from=date(26, 9, 1))) == "invalid_valid_from"
    assert _code(add(valid_from=None)) == "valid_from_required"
    assert _code(add(employee=None)) == "employee_required"
    top = _wage(manager, person, "100000000", date(2026, 2, 1))
    services.remove_wage(manager, top, "Deneme")
    created = AuditLog.objects.get(
        content_type__model="employeewage", object_id=wage.pk, action=AuditLog.Action.CREATE
    )
    assert created.actor == manager
    assert _code(lambda: services.remove_wage(manager, wage, " ")) == "reason_required"
    removed = services.remove_wage(manager, wage, "Yanlış tutar")
    assert not removed.is_active and removed.removed_by == manager
    assert AuditLog.objects.filter(
        content_type__model="employeewage",
        object_id=wage.pk,
        action=AuditLog.Action.UPDATE,
        note="Ücret kaydı kaldırıldı",
    ).exists()
    assert _code(lambda: services.remove_wage(manager, removed, "Tekrar")) == "already_removed"
    # Kaldırılan ücret hesaba ve geçmişe girmez.
    assert _row(2026, 3, person).missing_wage
    assert selectors.wage_summary(date(2026, 3, 1))[0]["history"] == []
    _wage(manager, person, "31000", date(2026, 1, 1))  # kaldırılan tarihe yeniden girilebilir
    summary = selectors.wage_summary(date(2026, 3, 1))
    assert [row["current_wage"] for row in summary] == [D("31000.00")]
    assert [item.gross_monthly for item in summary[0]["history"]] == [D("31000.00")]
    assert _row(2026, 3, person).gross_salary == D("31000.00")

    for rate in (D("-1"), D("100.01")):
        code = _code(partial(services.update_personnel_cost_settings, manager, rate))
        assert code == "invalid_employer_rate"
    # Okumak ayar kaydı açmaz; ilk değişiklik kaydı denetimle açar.
    assert services.get_personnel_cost_settings().employer_rate == D("17.75")
    assert not PersonnelCostSettings.objects.exists()
    settings = services.update_personnel_cost_settings(manager, D("20"))
    assert settings.created_by == manager and PersonnelCostSettings.objects.count() == 1
    assert AuditLog.objects.filter(
        content_type__model="personnelcostsettings", action=AuditLog.Action.CREATE
    ).exists()
    services.update_personnel_cost_settings(manager, D("17.75"))
    assert PersonnelCostSettings.objects.get().employer_rate == D("17.75")
    assert (
        _code(lambda: services.set_department_cost_center(manager, "Dolum", "SALES"))
        == "invalid_stage"
    )


def test_wage_summary_lists_leavers_on_request(
    client_for: Callable[..., APIClient],
    manager: User,
    django_assert_max_num_queries: Callable[..., Any],
) -> None:
    today = timezone.localdate()
    working = employee(manager)
    leaving = employee(manager, full_name="Can Kaya")
    left = employee(manager, full_name="Ece Ak", hire_date=date(2020, 1, 1))
    personnel.terminate_employee(manager, leaving, today + timedelta(days=10))
    personnel.terminate_employee(manager, left, today - timedelta(days=40))
    names = [row["employee"].full_name for row in selectors.wage_summary()]
    assert names == ["Ayşe Demir", "Can Kaya"]  # çıkışı ileride olan listede kalır
    names = [row["employee"].full_name for row in selectors.wage_summary(include_former=True)]
    assert sorted(names) == ["Ayşe Demir", "Can Kaya", "Ece Ak"]
    client = client_for(YONETIM)
    rows = client.get(f"{URL}/wages/summary/", {"include_former": "true"}).json()
    ece = next(row for row in rows if row["employee"]["full_name"] == "Ece Ak")
    assert ece["employee"]["is_active"] is False
    assert ece["employee"]["termination_date"] == str(today - timedelta(days=40))
    # Ayrılan kişiye ücret girilebilir (raporda eksik görünen geçmiş ay için).
    _wage(manager, left, "30000", date(2020, 1, 1))
    _wage(manager, left, "32000", date(2021, 1, 1))
    _wage(manager, working, "30000", date(2020, 1, 1))
    assert working.is_active
    # Geçmiş ücretler kişi başına ek sorgu açmaz.
    with django_assert_max_num_queries(3):
        data = WageSummarySerializer(selectors.wage_summary(include_former=True), many=True).data
    assert [len(row["history"]) for row in data if row["employee"]["full_name"] == "Ece Ak"] == [2]


def test_concurrent_duplicate_wage_is_a_business_error(
    manager: User, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Aynı kişi ve tarihe aynı anda iki ücret: tekil kısıt 500 değil `duplicate_wage` verir."""
    person = employee(manager)

    def clash(*args: Any, **kwargs: Any) -> Any:
        raise IntegrityError("costing_wage_employee_date_uniq")

    monkeypatch.setattr(services, "save_with_audit", clash)
    assert _code(lambda: _wage(manager, person, "30000", date(2026, 1, 1))) == "duplicate_wage"


@pytest.mark.parametrize("role", OTHER_ROLES)
def test_only_management_handles_wages(
    make_user: Callable[..., User], manager: User, role: str
) -> None:
    person = employee(manager)
    wage = _wage(manager, person, "30000", date(2026, 1, 1))
    other = make_user(role)
    calls = [
        lambda: _wage(other, person, "30000", date(2026, 2, 1)),
        lambda: services.remove_wage(other, wage, "Neden"),
        lambda: services.update_personnel_cost_settings(other, D(10)),
        lambda: services.set_department_cost_center(other, "Dolum", CostStage.FILLING),
        lambda: services.add_labor_rate_from_personnel_cost(
            other, 2026, 9, CostStage.PRODUCTION, date(2026, 10, 1)
        ),
    ]
    for call in calls:
        with pytest.raises(PermissionDenied):
            call()


# --- API -------------------------------------------------------------------------------------


def test_report_wages_centers_and_settings_api(
    client_for: Callable[..., APIClient], manager: User
) -> None:
    client = client_for(YONETIM)
    person = employee(manager, employee_no="S-1", department="Dolum")
    response = client.post(
        f"{URL}/wages/",
        {"employee": person.pk, "gross_monthly": "30000", "valid_from": "2020-01-01"},
        format="json",
    )
    assert response.status_code == 201, response.json()
    wage = response.json()
    assert wage["employee_name"] == "Ayşe Demir" and wage["gross_monthly"] == "30000.00"
    response = client.post(
        f"{URL}/wages/",
        {"employee": person.pk, "gross_monthly": "1", "valid_from": "2020-01-01"},
        format="json",
    )
    assert response.status_code == 400 and response.json()["code"] == "duplicate_wage"
    summary = client.get(f"{URL}/wages/summary/").json()
    assert summary[0]["current_wage"] == "30000.00" and summary[0]["stage"] == ""
    assert [item["id"] for item in summary[0]["history"]] == [wage["id"]]
    assert client.delete(f"{URL}/wages/{wage['id']}/").status_code == 405

    response = client.put(
        f"{URL}/cost-centers/", {"department": "dolum", "stage": "FILLING"}, format="json"
    )
    assert response.status_code == 200, response.json()
    assert client.get(f"{URL}/cost-centers/").json() == [
        {"department": "Dolum", "stage": "FILLING"}
    ]
    settings = client.get(f"{URL}/personnel-cost-settings/").json()
    assert settings == {"employer_rate": "17.75", "updated_at": None}
    response = client.patch(
        f"{URL}/personnel-cost-settings/", {"employer_rate": "101"}, format="json"
    )
    assert response.status_code == 400 and response.json()["code"] == "invalid_employer_rate"
    response = client.patch(f"{URL}/personnel-cost-settings/", {}, format="json")
    assert response.status_code == 400 and response.json()["code"] == "validation_error"

    report = client.get(f"{URL}/personnel-costs/2026/9/").json()
    assert report["period_label"] == "Eylül 2026" and report["attendance_approved"] is False
    assert report["total_cost"] == "35325.00" and report["employer_rate"] == "17.75"
    row = report["rows"][0]
    assert row["employee"]["employee_no"] == "S-1" and row["stage"] == "FILLING"
    assert row["paid_days"] == 30 and row["gross_total"] == "30000.00"
    stages = {item["stage"]: item for item in report["stages"]}
    assert stages["FILLING"]["cost"] == "35325.00" and stages["FILLING"]["suggested_rate"] is None
    assert client.get(f"{URL}/personnel-costs/2026/13/").status_code == 404

    excel = client.get(f"{URL}/personnel-costs/2026/9/xlsx/")
    assert excel.status_code == 200
    assert excel["Content-Disposition"].startswith(
        'attachment; filename="personel-maliyeti-2026-09-'
    )
    workbook = load_workbook(BytesIO(excel.content))
    assert workbook.sheetnames == ["Personel maliyeti", "Aşamalar", "Özet"]
    (person_row,) = _sheet_rows(workbook["Personel maliyeti"])
    assert person_row["Sicil no"] == "S-1" and person_row["Ad soyad"] == "Ayşe Demir"
    assert person_row["Maliyet merkezi"] == "Dolum (mamul)"
    assert person_row["Ödenen gün"] == 30 and person_row["Ücret eksik"] in ("", None)
    assert D(str(person_row["Brüt ücret (TL)"])) == D("30000")
    assert D(str(person_row["İşveren payı (TL)"])) == D("5325")
    assert D(str(person_row["Toplam maliyet (TL)"])) == D("35325")
    stage_rows = {row["Maliyet merkezi"]: row for row in _sheet_rows(workbook["Aşamalar"])}
    assert list(stage_rows) == ["Üretim (bulk)", "Dolum (mamul)", "Genel gider"]
    assert D(str(stage_rows["Dolum (mamul)"]["Personel maliyeti (TL)"])) == D("35325")
    assert stage_rows["Genel gider"]["Kişi"] == 0
    assert workbook["Özet"]["B3"].value == "Açık (taslak)"

    response = client.post(
        f"{URL}/personnel-costs/2026/9/labor-rate/",
        {"stage": "FILLING", "valid_from": "2026-10-01"},
        format="json",
    )
    assert response.status_code == 400 and response.json()["code"] == "attendance_not_approved"
    response = client.post(f"{URL}/wages/{wage['id']}/remove/", {"reason": "Yanlış"})
    assert response.status_code == 200 and response.json()["is_active"] is False


def test_admin_reaches_personnel_costs_and_anonymous_is_rejected(
    client_for: Callable[..., APIClient], api_client: APIClient, manager: User
) -> None:
    person = employee(manager)
    admin = client_for(ADMIN)
    response = admin.post(
        f"{URL}/wages/",
        {"employee": person.pk, "gross_monthly": "30000", "valid_from": "2020-01-01"},
        format="json",
    )
    assert response.status_code == 201
    assert admin.get(f"{URL}/personnel-costs/2026/9/").status_code == 200
    assert admin.get(f"{URL}/personnel-costs/2026/9/xlsx/").status_code == 200
    assert admin.get(f"{URL}/wages/summary/").status_code == 200
    response = admin.put(
        f"{URL}/cost-centers/", {"department": "Dolum", "stage": "FILLING"}, format="json"
    )
    assert response.status_code == 200
    response = admin.patch(
        f"{URL}/personnel-cost-settings/", {"employer_rate": "18"}, format="json"
    )
    assert response.status_code == 200
    for url in (
        f"{URL}/personnel-costs/2026/9/",
        f"{URL}/wages/summary/",
        f"{URL}/cost-centers/",
        f"{URL}/personnel-cost-settings/",
    ):
        assert api_client.get(url).status_code in (401, 403)


@pytest.mark.parametrize("role", OTHER_ROLES)
def test_other_roles_cannot_reach_personnel_costs(
    client_for: Callable[..., APIClient], manager: User, role: str
) -> None:
    person = employee(manager)
    wage = _wage(manager, person, "30000", date(2026, 1, 1))
    client = client_for(role)
    checks = [
        client.get(f"{URL}/personnel-costs/2026/9/"),
        client.get(f"{URL}/personnel-costs/2026/9/xlsx/"),
        client.post(
            f"{URL}/personnel-costs/2026/9/labor-rate/",
            {"stage": "PRODUCTION", "valid_from": "2026-10-01"},
        ),
        client.get(f"{URL}/wages/"),
        client.get(f"{URL}/wages/summary/"),
        client.post(f"{URL}/wages/", {"employee": person.pk, "gross_monthly": "1"}),
        client.post(f"{URL}/wages/{wage.pk}/remove/", {"reason": "X"}),
        client.get(f"{URL}/cost-centers/"),
        client.put(f"{URL}/cost-centers/", {"department": "Dolum", "stage": "FILLING"}),
        client.get(f"{URL}/personnel-cost-settings/"),
        client.patch(f"{URL}/personnel-cost-settings/", {"employer_rate": "10"}),
    ]
    assert [response.status_code for response in checks] == [403] * len(checks)
    assert EmployeeWage.objects.count() == 1
