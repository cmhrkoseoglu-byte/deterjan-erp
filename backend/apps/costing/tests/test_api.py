from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.core.roles import ADMIN, DEPO, LABORATUVAR, PLANLAMA, SATIN_ALMA, URETIM, YONETIM
from apps.costing.models import LotCost
from apps.costing.tests.factories import build_chain
from apps.inventory import services as inventory
from apps.masterdata.tests.factories import MaterialFactory, SupplierFactory, unit
from apps.production.tests.factories import build_plant

D = Decimal
URL = "/api/v1/costing"
REPORTS = ("cost-trend", "batch-variance", "scrap-cost", "top-materials", "stock-value")
OTHER_ROLES = (DEPO, PLANLAMA, URETIM, LABORATUVAR, SATIN_ALMA)

pytestmark = pytest.mark.usefixtures("invariants")


@pytest.fixture
def chain(staff, commit):
    return build_chain(staff, commit)


def test_lot_costs_list_detail_and_history(client_for, chain):
    client = client_for(YONETIM)
    response = client.get(f"{URL}/lot-costs/", {"method": "FILLING"})
    assert response.status_code == 200
    (row,) = response.json()["results"]
    assert row["lot_no"] == chain.run.output_lot.lot_no
    assert row["unit_cost"] == "37.2019"
    detail = client.get(f"{URL}/lot-costs/{row['id']}/").json()
    assert [line["category"] for line in detail["lines"]][:2] == ["BULK", "PACKAGING"]
    assert detail["finished_scrap_cost"] == "366.12"
    history = client.get(f"{URL}/lots/{chain.batch.output_lot.pk}/").json()
    assert [v["version"] for v in history["versions"]] == [1]
    assert history["is_costable"] is True


def test_recalculate_api(client_for, chain):
    client = client_for(YONETIM)
    lot = chain.raw["LAS"]
    response = client.post(f"{URL}/lots/{lot.pk}/recalculate/", {"note": ""}, format="json")
    assert response.status_code == 400
    response = client.post(
        f"{URL}/lots/{lot.pk}/recalculate/", {"note": "Kontrol", "cascade": True}, format="json"
    )
    assert response.status_code == 200
    assert [v["version"] for v in response.json()["versions"]] == [2, 1]
    assert LotCost.objects.get(lot=chain.run.output_lot, is_current=True).version == 2


def test_rates_and_standard_costs_api(client_for, db):
    client = client_for(YONETIM)
    data = {
        "stage": "PRODUCTION",
        "kind": "ENERGY",
        "rate_per_kg": "0.75",
        "valid_from": "2026-01-01",
    }
    response = client.post(f"{URL}/rates/", data, format="json")
    assert response.status_code == 201, response.json()
    rate_id = response.json()["id"]
    duplicate = client.post(f"{URL}/rates/", data, format="json")
    assert duplicate.status_code == 400
    assert duplicate.json()["code"] == "duplicate_rate"
    response = client.patch(f"{URL}/rates/{rate_id}/", {"rate_per_kg": "0.80"}, format="json")
    assert response.json()["rate_per_kg"] == "0.8000"
    assert client.delete(f"{URL}/rates/{rate_id}/").status_code == 405

    material = MaterialFactory.create()
    response = client.post(
        f"{URL}/standard-costs/", {"material": material.pk, "unit_cost": "12.5"}, format="json"
    )
    assert response.status_code == 201
    cost_id = response.json()["id"]
    response = client.patch(f"{URL}/standard-costs/{cost_id}/", {"is_active": False}, format="json")
    assert response.json()["is_active"] is False


def test_reports_and_excel(client_for, chain):
    client = client_for(YONETIM)
    trend = client.get(f"{URL}/reports/cost-trend/").json()
    assert {r["material_code"] for r in trend} == {"BULASIK", "BULASIK-750"}
    variance = client.get(f"{URL}/reports/batch-variance/").json()
    (row,) = variance
    # Teorik: 600 × 46,0618 + 300 × 40,9619 + 100 × 0,1829 = 39.943,94 (aynı fiyatlar).
    assert row["theoretical_material_cost"] == "39943.94"
    assert row["actual_material_cost"] == "39943.94"
    assert row["theoretical_unit_cost"] == "39.9439"
    assert row["actual_material_unit_cost"] == "40.7591"
    assert row["variance_pct"] == "2.04"
    (scrap,) = client.get(f"{URL}/reports/scrap-cost/").json()
    assert scrap["material_scrap_cost"] == "223.80"
    assert scrap["total_scrap_cost"] == "589.92"
    top = client.get(f"{URL}/reports/top-materials/").json()
    assert [r["material_code"] for r in top[:2]] == ["LAS", "SLES"]
    assert top[0]["amount"] == "27637.08"
    assert (
        client.get(f"{URL}/reports/top-materials/", {"material_type": "PACKAGING"}).json()[0][
            "material_code"
        ]
        == "SISE-750"
    )
    value = client.get(f"{URL}/reports/stock-value/").json()
    las = next(r for r in value["rows"] if r["material_code"] == "LAS")
    # 2.000 − 600 = 1.400 kg × 46,0618.
    assert (las["quantity"], las["value"]) == ("1400.0000", "64486.52")
    assert D(value["total_value"]) == sum(D(r["value"]) for r in value["rows"])
    for name in REPORTS:
        response = client.get(f"{URL}/reports/{name}/xlsx/")
        assert response.status_code == 200, name
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.active.max_row >= 2, name


def test_report_query_validation(client_for, db):
    client = client_for(YONETIM)
    assert client.get(f"{URL}/reports/top-materials/", {"limit": 0}).status_code == 400
    assert client.get(f"{URL}/reports/cost-trend/", {"date_from": "x"}).status_code == 400


@pytest.mark.parametrize("role", OTHER_ROLES)
def test_cost_amounts_are_forbidden_to_other_roles(client_for, db, role):
    client = client_for(role)
    paths = [
        f"{URL}/lot-costs/",
        f"{URL}/rates/",
        f"{URL}/standard-costs/",
        f"{URL}/lots/1/",
        *(f"{URL}/reports/{name}/" for name in REPORTS),
        *(f"{URL}/reports/{name}/xlsx/" for name in REPORTS),
    ]
    for path in paths:
        assert client.get(path).status_code == 403, path
    assert client.post(f"{URL}/lots/1/recalculate/", {"note": "x"}).status_code == 403
    assert client.post(f"{URL}/rates/", {}).status_code == 403


def test_admin_sees_costs_and_anonymous_is_rejected(client_for, api_client, db):
    assert client_for(ADMIN).get(f"{URL}/lot-costs/").status_code == 200
    assert api_client.get(f"{URL}/lot-costs/").status_code in (401, 403)


def _receipt(staff):
    plant = build_plant(staff[PLANLAMA], stock=False)
    receipt = inventory.create_goods_receipt(
        staff[DEPO],
        {
            "supplier": SupplierFactory.create(),
            "receipt_date": date.today(),
            "quarantine_location": plant.quarantine,
        },
        [{"material": plant.las, "quantity": D("10"), "unit": unit("KG"), "unit_price": D("5")}],
    )
    return receipt


def test_receipt_prices_hidden_from_depo_and_ignored_on_write(client_for, staff):
    receipt = _receipt(staff)
    line = receipt.lines.get()
    depo = client_for(DEPO)
    body = depo.get(f"/api/v1/inventory/goods-receipts/{receipt.pk}/").json()
    assert "unit_price" not in body["lines"][0]
    assert body["extra_costs"] is None
    response = depo.patch(
        f"/api/v1/inventory/goods-receipts/{receipt.pk}/",
        {
            "lines": [
                {
                    "id": line.pk,
                    "material": line.material_id,
                    "quantity": "12",
                    "unit": line.unit_id,
                    "unit_price": "999",
                }
            ]
        },
        format="json",
    )
    assert response.status_code == 200, response.json()
    line.refresh_from_db()
    assert (line.quantity, line.unit_price) == (D("12"), D("5"))
    assert depo.post(f"/api/v1/inventory/goods-receipts/{receipt.pk}/costs/", {}).status_code == 403


@pytest.mark.parametrize("role", (SATIN_ALMA, YONETIM))
def test_receipt_costs_endpoint_for_price_roles(client_for, staff, role):
    receipt = _receipt(staff)
    line = receipt.lines.get()
    client = client_for(role)
    response = client.post(
        f"/api/v1/inventory/goods-receipts/{receipt.pk}/costs/",
        {
            "lines": [{"id": line.pk, "unit_price": "2", "currency": "EUR", "exchange_rate": "35"}],
            "extra_costs": [{"kind": "FREIGHT", "amount": "50", "allocation": "VALUE"}],
        },
        format="json",
    )
    assert response.status_code == 200, response.json()
    body = response.json()
    assert body["lines"][0]["currency"] == "EUR"
    assert body["lines"][0]["exchange_rate"] == "35.000000"
    assert body["extra_costs"][0]["amount"] == "50.00"
    response = client.post(
        f"/api/v1/inventory/goods-receipts/{receipt.pk}/costs/",
        {"lines": [{"id": line.pk, "currency": "USD", "exchange_rate": None}]},
        format="json",
    )
    assert response.status_code == 400
    assert response.json()["code"] == "exchange_rate_required"
